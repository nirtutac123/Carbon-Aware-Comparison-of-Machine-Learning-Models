from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "PES-Technical-Report-Template_Jan_2019.docx"
REPORT_DIR = ROOT / "report"
OUT = REPORT_DIR / "Carbon_Aware_ML_Germany_IEEEPES_Technical_Report.docx"
FIG_DIR = ROOT / "outputs" / "figures"
TABLE_DIR = ROOT / "outputs" / "tables"


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc: Document, text: str = "", style: str | None = None, align=None, size=10.5):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run(r, size=size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run(r, size=13 if level == 1 else 11.5, bold=True, color=(31, 43, 56))
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold=False, size=8.5, color=(20, 30, 40), align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc: Document, text: str, table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=9.5 if table else 9.2, bold=True)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_df_table(doc: Document, caption: str, df: pd.DataFrame, columns: list[tuple[str, str]], max_rows=None):
    add_caption(doc, caption, table=True)
    if max_rows is not None:
        df = df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = True
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for j, (_, label) in enumerate(columns):
        cell = table.rows[0].cells[j]
        shade_cell(cell, "DCEAF7")
        set_cell_text(cell, label, bold=True, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, (col, _) in enumerate(columns):
            value = row[col]
            if isinstance(value, float):
                if "time" in col or "carbon" in col or "size" in col or "energy" in col:
                    value = f"{value:.3f}"
                else:
                    value = f"{value:.3f}"
            set_cell_text(cells[j], value, size=7.8, align=WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT)
    add_para(doc, "", size=4)
    return table


def add_picture(doc: Document, filename: str, caption: str, width=6.55):
    path = FIG_DIR / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("title", caption.split(". ", 1)[0])
    inline_shape._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.3)
        if section.header.paragraphs:
            section.header.paragraphs[0].text = "PS26 Technical Report | Carbon-Aware Machine Learning"
            for run in section.header.paragraphs[0].runs:
                set_run(run, size=8.5, color=(90, 90, 90))
        if section.footer.paragraphs:
            section.footer.paragraphs[0].text = "IEEE/PES-style course report | May 2026"
            section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in section.footer.paragraphs[0].runs:
                set_run(run, size=8.5, color=(90, 90, 90))
    for style_name in ["Normal", "Body Text"]:
        if style_name in [s.name for s in doc.styles]:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(10.5)


def pct(x):
    return f"{100*x:.1f}%"


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_document(doc)

    rq = pd.read_csv(TABLE_DIR / "research_questions.csv")
    rq1 = pd.read_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv")
    rq2 = pd.read_csv(TABLE_DIR / "rq2_feature_ablation_metrics.csv")
    rq3 = pd.read_csv(TABLE_DIR / "rq3_seasonal_generalization.csv")
    rq4 = pd.read_csv(TABLE_DIR / "rq4_top_interpretable_features.csv")
    rq5 = pd.read_csv(TABLE_DIR / "rq5_training_window_scenarios.csv")
    metadata = json.loads((TABLE_DIR / "dataset_metadata.json").read_text(encoding="utf-8"))

    best = rq1.sort_values("f1_score", ascending=False).iloc[0]
    low_compute = rq1.sort_values("training_carbon_mgco2e_proxy").iloc[0]
    hgb_rq5 = rq5[rq5["model"] == best["model"]]
    green_reduction = hgb_rq5[hgb_rq5["scenario"].str.startswith("Green")]["reduction_vs_random_percent"].iloc[0]

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ACCURACY VERSUS SUSTAINABILITY")
    set_run(r, size=20, bold=True, color=(31, 43, 56))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A Carbon-Aware Comparison of Machine Learning Models for German Power-System Load Risk")
    set_run(r, size=15, bold=True, color=(31, 43, 56))
    add_para(doc, "", size=4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IEEE Power & Energy Society style technical report")
    set_run(r, size=12, italic=True, color=(60, 70, 80))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared for PS26 - Interdisciplinary Elective: AI, Power & Responsibility")
    set_run(r, size=11.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared by: PS26 Project Team | Date: 03 May 2026")
    set_run(r, size=10.5)
    add_para(doc, "", size=8)
    add_picture(doc, "workflow_structured_ai_esg_style.png", "Fig. 1. Structured AI workflow for carbon-aware sustainability auditing.", width=6.6)
    doc.add_page_break()

    add_heading(doc, "ACKNOWLEDGMENTS", 1)
    add_para(doc, "This report uses the Open Power System Data time-series package and the IEEE PES technical report template supplied with the course prompt. The analysis was generated as a reproducible research package with raw data, processed data, notebooks, publication figures, CSV tables, reusable vector icons, and editable poster assets.")

    add_heading(doc, "KEYWORDS", 1)
    add_para(doc, "carbon-aware machine learning; Germany; high-load prediction; IEEE PES; model selection; power systems; renewable generation; responsible AI; sustainability; technical reports")

    add_heading(doc, "CONTENTS", 1)
    contents = [
        "Executive Summary",
        "1. Introduction",
        "2. Dataset, Research Questions, and Scope",
        "3. Methodology",
        "4. Results",
        "5. Responsible AI and Power-System Governance Implications",
        "6. Limitations",
        "7. Conclusions and Recommendations",
        "Appendix A. Reproducibility Package",
        "References",
    ]
    for item in contents:
        add_para(doc, item, size=10.5)
    doc.add_page_break()

    add_heading(doc, "EXECUTIVE SUMMARY", 1)
    add_para(doc, f"This technical report evaluates carbon-aware machine-learning model selection for Germany-aligned power-system analytics. The empirical task is next-day high-load-event classification: given hourly operational signals at time t, predict whether German load at t+24 hours is in the highest quintile of observed demand. The final processed dataset contains {metadata['n_rows']:,} hourly observations and {metadata['n_features']} engineered features spanning calendar cycles, load history, wind and solar generation, residual-load proxies, TSO-zone state, and day-ahead load forecasts.")
    add_para(doc, f"The best overall predictive model in the generated experiment is {best['model']}, with F1 = {best['f1_score']:.3f}, ROC-AUC = {best['roc_auc']:.3f}, and an estimated training carbon proxy of {best['training_carbon_mgco2e_proxy']:.2f} mg CO2e. The lowest-compute baseline is {low_compute['model']}, with a much smaller carbon proxy ({low_compute['training_carbon_mgco2e_proxy']:.2f} mg CO2e) but lower F1 ({low_compute['f1_score']:.3f}). A carbon-aware training-window simulation indicates that scheduling the selected model into the lowest residual-load decile reduces the training carbon proxy by {green_reduction:.1f}% relative to random-hour scheduling.")
    add_para(doc, "The central conclusion is not that a single model is universally best. Rather, responsible model selection should be Pareto-based: select the simplest model whose marginal improvement justifies its marginal computational, carbon, and auditability costs. For this course-scale study, histogram gradient boosting is the strongest recommendation for balanced accuracy and sustainability; linear SVM or logistic regression remain defensible for ultra-low-compute monitoring contexts.")

    add_heading(doc, "1. INTRODUCTION", 1)
    add_para(doc, "Machine-learning systems increasingly mediate decisions in energy operations, forecasting, asset management, and grid planning. In the power sector, predictive accuracy is important, but model choice also consumes computation, imposes maintenance burdens, and can conflict with sustainability goals. This makes power-system machine learning a natural test case for the PS26 theme: AI, power, and responsibility.")
    add_para(doc, "IEEE PES technical reports and white papers are designed for deeper treatment than short conference papers. The IEEE PES technical report page describes these reports as vehicles for state-of-technical-research synthesis, background for possible future standards or guides, emerging-technology analysis, surveys, strategic roadmaps, and technical or policy positions. Recent PES Resource Center examples follow a committee-report pattern: a clear scope, long-form technical context, methodology, evidence, gaps, and recommendations.")
    add_para(doc, "This report adapts that pattern to a course-scale empirical project. It provides an end-to-end reproducible workflow, not merely a narrative summary: each research question has an executable notebook, one CSV result table, and one publication-ready PDF figure.")

    add_heading(doc, "2. DATASET, RESEARCH QUESTIONS, AND SCOPE", 1)
    add_heading(doc, "2.1 Dataset Selection", 2)
    add_para(doc, "The selected public dataset is the Open Power System Data time-series package, version 2020-10-06. The package provides load, wind and solar generation, and prices in hourly resolution, aggregated by country, control area, or bidding zone. For Germany, the field documentation includes national load, day-ahead load forecast, solar generation, wind generation, and TSO control-area load variables. This directly matches the project topic and avoids using a generic benchmark detached from German power-system conditions.")
    dataset_table = pd.DataFrame(
        [
            ["Raw file", "data.csv", "OPSD 60-minute single-index CSV"],
            ["Processed file", "preprocessed_content.csv", "Germany-specific engineered ML table"],
            ["Date range", "2015-01 to 2020-09", metadata["date_range_utc"][0][:10] + " to " + metadata["date_range_utc"][1][:10]],
            ["Target", "Next-day high load", metadata["target_definition"]],
            ["Train/test split", "Chronological", "2015-2018 training; 2019-2020 test"],
            ["Carbon variable", "Residual-load proxy", "150 to 650 gCO2e/kWh proxy, not measured grid emissions"],
        ],
        columns=["item", "value", "detail"],
    )
    add_df_table(doc, "TABLE 1. Dataset and target definition.", dataset_table, [("item", "Item"), ("value", "Value"), ("detail", "Detail")])

    add_heading(doc, "2.2 Research Questions", 2)
    add_df_table(doc, "TABLE 2. Final research questions and generated artifacts.", rq, [("id", "ID"), ("question", "Research question"), ("figure", "Figure"), ("table", "Table")])

    add_heading(doc, "3. METHODOLOGY", 1)
    add_heading(doc, "3.1 Preprocessing and Feature Engineering", 2)
    add_para(doc, "The raw OPSD table is filtered to Germany and German TSO-zone variables. Sparse gaps are filled by time-ordered interpolation and forward/backward filling. Calendar variables are encoded cyclically. Load history is represented by 1-hour, 24-hour, and 168-hour lags, together with 24-hour and 168-hour rolling means. Renewable operating state is represented by wind, solar, aggregate wind-plus-solar generation, renewable share, and residual-load variables. Regional structure is represented by 50Hertz, Amprion, TenneT, and TransnetBW load values and a TSO-zone imbalance feature.")
    add_para(doc, "The target is explicitly predictive rather than contemporaneous: the feature vector at hour t predicts whether Germany's load at t+24 hours is in the highest quintile. This makes current load, current renewable generation, and current TSO-zone state valid predictors rather than target leakage.")

    add_heading(doc, "3.2 Model Suite and Evaluation", 2)
    add_para(doc, "The model suite compares logistic regression, linear support vector machine, random forest, histogram gradient boosting, and a shallow neural network. The gradient-boosting baseline is a reproducible histogram gradient boosting classifier from scikit-learn. It is used as an XGBoost-class boosted-tree comparator without adding an external XGBoost dependency to the submitted package.")
    add_para(doc, "All models are trained on 2015-2018 observations and evaluated on 2019-2020 observations. Metrics include accuracy, balanced accuracy, precision, recall, F1-score, ROC-AUC, fit time, prediction time, serialized model size, proxy energy, and proxy carbon. F1-score is emphasized because the high-load class is the operational event of interest.")

    add_heading(doc, "3.3 Carbon-Aware Accounting", 2)
    add_para(doc, "The report uses a transparent proxy rather than claiming measured emissions. Runtime energy is estimated from observed fit and prediction times using assumed CPU powers of 45 W for training and 18 W for inference. Proxy carbon is then calculated by multiplying energy by a residual-load carbon-intensity proxy. The residual-load proxy increases as the load share not served by observed wind and solar increases. This is suitable for ranking and sensitivity analysis, but it is not a substitute for hardware energy metering or official grid-emission factors.")

    doc.add_page_break()
    add_heading(doc, "4. RESULTS", 1)
    add_heading(doc, "4.1 RQ1: Accuracy-Carbon Pareto Frontier", 2)
    add_para(doc, f"{best['model']} is the strongest predictive model, while {low_compute['model']} is the lowest-compute baseline. The Pareto frontier therefore contains both high-accuracy and low-compute candidates. This is the main governance result: the preferred model depends on whether the deployment context values maximum event detection, minimum compute, or auditability.")
    add_df_table(
        doc,
        "TABLE 3. Model trade-off metrics for next-day high-load classification.",
        rq1[["model", "f1_score", "roc_auc", "fit_time_s", "model_size_kb", "training_carbon_mgco2e_proxy", "pareto_efficient"]],
        [
            ("model", "Model"),
            ("f1_score", "F1"),
            ("roc_auc", "ROC-AUC"),
            ("fit_time_s", "Fit time (s)"),
            ("model_size_kb", "Size (KB)"),
            ("training_carbon_mgco2e_proxy", "Carbon proxy (mg CO2e)"),
            ("pareto_efficient", "Pareto"),
        ],
    )
    add_picture(doc, "rq1_accuracy_carbon_pareto.png", "Fig. 2. Accuracy-carbon Pareto frontier for the five-model suite.")

    doc.add_page_break()
    add_heading(doc, "4.2 RQ2: Feature Value", 2)
    add_para(doc, "The feature-ablation experiment shows that calendar variables alone are surprisingly strong because German electricity demand has pronounced daily, weekly, and seasonal regularity. Adding national load history yields the largest marginal improvement. Renewable and TSO-zone variables add operational context, but the full feature set does not dominate the national-history feature set in this run, suggesting that parsimonious feature selection is part of carbon-aware modeling.")
    add_df_table(
        doc,
        "TABLE 4. Feature-ablation metrics using histogram gradient boosting.",
        rq2[["feature_set", "n_features", "f1_score", "roc_auc", "training_carbon_mgco2e_proxy"]],
        [
            ("feature_set", "Feature set"),
            ("n_features", "No. features"),
            ("f1_score", "F1"),
            ("roc_auc", "ROC-AUC"),
            ("training_carbon_mgco2e_proxy", "Carbon proxy (mg CO2e)"),
        ],
    )
    add_picture(doc, "rq2_feature_ablation.png", "Fig. 3. Feature value versus training carbon proxy.")

    doc.add_page_break()
    add_heading(doc, "4.3 RQ3: Seasonal Generalization", 2)
    seasonal_summary = rq3.pivot(index="model", columns="season", values="f1_score").reset_index()
    seasonal_summary = seasonal_summary[["model", "Winter", "Spring", "Summer", "Autumn"]]
    add_para(doc, "The seasonal stress test confirms that winter is the easiest regime and summer is the hardest, especially for the linear baselines. Neural and boosted-tree models maintain stronger summer robustness, while random forest remains competitive in autumn and winter. This pattern is consistent with the higher variance of summer renewable generation and demand conditions.")
    add_df_table(
        doc,
        "TABLE 5. Seasonal F1-score by model.",
        seasonal_summary,
        [("model", "Model"), ("Winter", "Winter"), ("Spring", "Spring"), ("Summer", "Summer"), ("Autumn", "Autumn")],
    )
    add_picture(doc, "rq3_seasonal_f1_heatmap.png", "Fig. 4. Seasonal F1-score heatmap on the 2019-2020 holdout period.")

    doc.add_page_break()
    add_heading(doc, "4.4 RQ4: Interpretability and Auditability", 2)
    add_para(doc, "The interpretable model identifies calendar structure, rolling load, TSO-zone load, and residual-load/renewable-share terms as prominent drivers. Because many power-system variables are collinear, coefficient signs must be interpreted as controlled statistical effects rather than causal claims. The permutation-importance check provides a second lens on realized predictive contribution.")
    add_df_table(
        doc,
        "TABLE 6. Top interpretable features from standardized logistic regression.",
        rq4[["display_feature", "standardized_logit_coefficient", "permutation_importance_mean", "direction"]].head(10),
        [
            ("display_feature", "Feature"),
            ("standardized_logit_coefficient", "Coefficient"),
            ("permutation_importance_mean", "Permutation importance"),
            ("direction", "Direction"),
        ],
    )
    add_picture(doc, "rq4_interpretability_efficiency.png", "Fig. 5. Interpretable drivers and accuracy-auditability trade-off.")

    doc.add_page_break()
    add_heading(doc, "4.5 RQ5: Carbon-Aware Training Windows", 2)
    rq5_short = rq5[rq5["model"].isin([best["model"], low_compute["model"]])][
        ["model", "scenario", "carbon_intensity_proxy_g_per_kwh", "training_carbon_mgco2e_proxy", "reduction_vs_random_percent"]
    ].copy()
    add_para(doc, f"For the selected model, green-window scheduling reduces proxy training carbon by {green_reduction:.1f}% relative to an average training hour. Although the absolute carbon values are small because the models are lightweight and the dataset is course-scale, the relative result is important: scheduling and energy-aware orchestration matter even before model architecture changes.")
    add_df_table(
        doc,
        "TABLE 7. Training-window scenarios for the selected and lowest-compute models.",
        rq5_short,
        [
            ("model", "Model"),
            ("scenario", "Scenario"),
            ("carbon_intensity_proxy_g_per_kwh", "Proxy intensity"),
            ("training_carbon_mgco2e_proxy", "Carbon proxy (mg CO2e)"),
            ("reduction_vs_random_percent", "Reduction vs random (%)"),
        ],
    )
    add_picture(doc, "rq5_carbon_aware_training_windows.png", "Fig. 6. Carbon-aware training-window sensitivity by model.")

    doc.add_page_break()
    add_heading(doc, "5. RESPONSIBLE AI AND POWER-SYSTEM GOVERNANCE IMPLICATIONS", 1)
    add_para(doc, "This study supports a procurement-style view of ML model selection. A model should not be approved solely because it has the highest accuracy. A responsible selection record should include the target definition, data provenance, chronological validation design, performance metrics by operating regime, model size, runtime, carbon-accounting assumption, and evidence that simpler alternatives were tested.")
    add_para(doc, "For German power-system applications, the strongest governance rule is seasonal stress testing. A model that performs well on average but fails in summer or transition seasons may be unsuitable for operational use. The second rule is traceable carbon accounting. Even a proxy makes hidden compute costs visible, and it creates a structure for replacing assumptions with measured hardware energy later.")

    add_heading(doc, "6. LIMITATIONS", 1)
    add_para(doc, "First, the carbon variable is a residual-load proxy, not measured German grid-emission intensity. It should be treated as a sensitivity indicator. Second, the analysis is limited to OPSD variables available in the hourly single-index package; adding weather, market, holiday, and industrial activity data would make the task more realistic. Third, serialized model size and CPU runtime are environment-specific. Fourth, the report uses histogram gradient boosting as a reproducible boosted-tree comparator, not the external XGBoost library. Fifth, feature coefficients are not causal estimates.")

    add_heading(doc, "7. CONCLUSIONS AND RECOMMENDATIONS", 1)
    add_para(doc, f"The recommended balanced model is {best['model']}, because it achieves the highest F1-score ({best['f1_score']:.3f}) with a modest carbon proxy and strong ROC-AUC ({best['roc_auc']:.3f}). The recommended low-compute fallback is {low_compute['model']}, which is appropriate when transparency, speed, and operational simplicity outweigh maximum predictive performance.")
    add_para(doc, "For a publication or conference presentation, the strongest story is: (1) define model selection as a multi-objective power-system governance problem; (2) use Germany-specific open data; (3) compare models on both accuracy and sustainability indicators; (4) stress-test results seasonally; and (5) recommend a Pareto-efficient model with explicit caveats.")

    add_heading(doc, "APPENDIX A. REPRODUCIBILITY PACKAGE", 1)
    add_para(doc, "The generated package contains the raw OPSD CSV (`data.csv`), processed modeling table (`preprocessed_content.csv`), five notebooks under `notebooks/`, result tables under `outputs/tables/`, PDF and PNG figures under `outputs/figures/`, reusable SVG icons under `icons/`, and an editable PowerPoint poster deck under `posters/`.")
    add_para(doc, "Run `python3 scripts/generate_analysis_assets.py` from the project root to regenerate the computational assets. The report itself is produced by `scripts/generate_report_docx.py`.")

    add_heading(doc, "REFERENCES", 1)
    refs = [
        "IEEE Power & Energy Society, Technical Reports & White Papers, https://ieee-pes.org/technical-activities/technical-reports-white-papers/.",
        "Open Power System Data, Data Package Time series, Version 2020-10-06, https://doi.org/10.25832/time_series/2020-10-06.",
        "Open Power System Data, Time series field documentation, https://data.open-power-system-data.org/time_series/.",
        "A. M. Stankovic et al., Methods for Analysis and Quantification of Power System Resilience, IEEE PES Technical Report TR 108, 2023.",
        "H. Chen et al., IEEE Power and Energy Technology Assessment and Roadmap, IEEE PES Technical Report TR 123, 2024.",
        "R. Jeffers et al., Enabling Climate Adaptation and Mitigation Through Grid Modernization, IEEE PES Technical Report TR 124, 2024.",
        "F. Pedregosa et al., Scikit-learn: Machine Learning in Python, Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
        "R. Schwartz, J. Dodge, N. A. Smith, and O. Etzioni, Green AI, Communications of the ACM, vol. 63, no. 12, pp. 54-63, 2020.",
    ]
    for i, ref in enumerate(refs, start=1):
        add_para(doc, f"[{i}] {ref}", size=9.5)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
