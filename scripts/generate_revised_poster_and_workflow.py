from __future__ import annotations

import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "outputs" / "figures"
TABLE_DIR = ROOT / "outputs" / "tables"
POSTER_DIR = ROOT / "posters"
PREVIEW_DIR = ROOT / "qa_render" / "posters_v2_source_previews"
OUT_PPTX = POSTER_DIR / "carbon_aware_ml_academic_posters_v2.pptx"


W, H = 48.0, 27.0


PALETTES = {
    "white": {
        "bg": (248, 252, 250),
        "header": (255, 255, 255),
        "panel": (255, 255, 255),
        "line": (205, 219, 226),
        "dark": (14, 45, 64),
        "teal": (0, 137, 123),
        "green": (47, 158, 68),
        "blue": (42, 111, 151),
        "muted": (89, 105, 116),
        "soft": (229, 246, 242),
    },
    "dark": {
        "bg": (242, 248, 247),
        "header": (2, 32, 52),
        "panel": (255, 255, 255),
        "line": (190, 209, 214),
        "dark": (8, 31, 44),
        "teal": (14, 141, 130),
        "green": (70, 174, 94),
        "blue": (35, 103, 146),
        "muted": (81, 96, 108),
        "soft": (225, 244, 239),
    },
    "minimal": {
        "bg": (255, 255, 255),
        "header": (245, 250, 249),
        "panel": (255, 255, 255),
        "line": (216, 226, 231),
        "dark": (19, 48, 67),
        "teal": (0, 128, 128),
        "green": (42, 157, 103),
        "blue": (48, 100, 162),
        "muted": (87, 96, 106),
        "soft": (237, 248, 245),
    },
}


def rgb(value):
    return RGBColor(*value)


def add_text(slide, x, y, w, h, text, size=14, bold=False, color=(20, 30, 40), align=None, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align if align is not None else PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_panel(slide, x, y, w, h, title, number, pal, fill=None, title_fill=None):
    fill = fill or pal["panel"]
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(pal["line"])
    shape.line.width = Pt(1.3)
    bar_h = 0.48
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(bar_h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(title_fill or pal["teal"])
    bar.line.fill.background()
    add_text(slide, x + 0.22, y + 0.04, 0.55, 0.36, f"{number:02d}", size=17, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.9, y + 0.07, w - 1.1, 0.34, title.upper(), size=18.5, bold=True, color=(255, 255, 255))
    return shape


def add_small_circle_label(slide, x, y, text, pal, fill=None, size=0.44, color=(255, 255, 255)):
    fill = fill or pal["green"]
    circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    circ.fill.solid()
    circ.fill.fore_color.rgb = rgb(fill)
    circ.line.color.rgb = rgb((255, 255, 255))
    circ.line.width = Pt(1)
    add_text(slide, x, y + size * 0.09, size, size * 0.58, text, size=13.5, bold=True, color=color, align=PP_ALIGN.CENTER)


def add_icon(slide, x, y, kind, pal, scale=1.0, color=None):
    color = color or pal["teal"]
    base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(0.78 * scale), Inches(0.78 * scale))
    base.fill.solid()
    base.fill.fore_color.rgb = rgb((239, 249, 247))
    base.line.color.rgb = rgb(color)
    base.line.width = Pt(1.5)
    cx = x + 0.39 * scale
    cy = y + 0.39 * scale
    if kind == "ai":
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + 0.22 * scale), Inches(y + 0.22 * scale), Inches(0.34 * scale), Inches(0.34 * scale))
        chip.fill.solid()
        chip.fill.fore_color.rgb = rgb((218, 244, 238))
        chip.line.color.rgb = rgb(color)
        for dx in [0.16, 0.61]:
            ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + dx * scale), Inches(cy), Inches(x + (0.22 if dx < 0.4 else 0.56) * scale), Inches(cy))
            ln.line.color.rgb = rgb(color)
            ln.line.width = Pt(1)
        add_text(slide, x + 0.25 * scale, y + 0.29 * scale, 0.28 * scale, 0.14 * scale, "AI", size=5.5 * scale, bold=True, color=color, align=PP_ALIGN.CENTER)
    elif kind == "leaf":
        leaf = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.25 * scale), Inches(y + 0.19 * scale), Inches(0.36 * scale), Inches(0.24 * scale))
        leaf.rotation = -28
        leaf.fill.solid()
        leaf.fill.fore_color.rgb = rgb((91, 184, 118))
        leaf.line.color.rgb = rgb(color)
    elif kind == "database":
        cyl = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CAN, Inches(x + 0.22 * scale), Inches(y + 0.20 * scale), Inches(0.36 * scale), Inches(0.42 * scale))
        cyl.fill.solid()
        cyl.fill.fore_color.rgb = rgb((219, 240, 249))
        cyl.line.color.rgb = rgb(color)
    elif kind == "audit":
        shield = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PENTAGON, Inches(x + 0.22 * scale), Inches(y + 0.18 * scale), Inches(0.36 * scale), Inches(0.43 * scale))
        shield.fill.solid()
        shield.fill.fore_color.rgb = rgb((232, 246, 240))
        shield.line.color.rgb = rgb(color)
        add_text(slide, x + 0.27 * scale, y + 0.30 * scale, 0.26 * scale, 0.12 * scale, "✓", size=8 * scale, bold=True, color=color, align=PP_ALIGN.CENTER)
    elif kind == "chart":
        for i, ht in enumerate([0.18, 0.30, 0.42]):
            b = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + (0.22 + i * 0.11) * scale), Inches(y + (0.60 - ht) * scale), Inches(0.07 * scale), Inches(ht * scale))
            b.fill.solid()
            b.fill.fore_color.rgb = rgb(color)
            b.line.fill.background()
    elif kind == "doc":
        d = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.FLOWCHART_DOCUMENT, Inches(x + 0.23 * scale), Inches(y + 0.18 * scale), Inches(0.34 * scale), Inches(0.45 * scale))
        d.fill.solid()
        d.fill.fore_color.rgb = rgb((248, 252, 252))
        d.line.color.rgb = rgb(color)
    return base


def add_bullets(slide, x, y, w, bullets, pal, icon_kind="leaf", size=10.8, gap=0.56):
    for i, bullet in enumerate(bullets):
        yy = y + i * gap
        add_icon(slide, x, yy - 0.03, icon_kind if i % 2 == 0 else "audit", pal, scale=0.48, color=pal["green"] if i % 2 == 0 else pal["blue"])
        add_text(slide, x + 0.55, yy, w - 0.55, 0.34, bullet, size=size, color=pal["dark"])


def add_metric(slide, x, y, w, label, value, pal, accent=None):
    accent = accent or pal["teal"]
    add_text(slide, x, y, w, 0.50, value, size=36, bold=True, color=accent, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + 0.62, w, 0.48, label, size=13.2, color=pal["muted"], align=PP_ALIGN.CENTER)


def add_rq_list(slide, x, y, w, pal):
    items = [
        ("RQ1", "Which model is Pareto-efficient?"),
        ("RQ2", "Which feature families add value?"),
        ("RQ3", "Are rankings seasonally stable?"),
        ("RQ4", "Can simple models explain drivers?"),
        ("RQ5", "Do green training windows help?"),
    ]
    for i, (tag, txt) in enumerate(items):
        yy = y + i * 0.82
        add_small_circle_label(slide, x, yy, tag.replace("RQ", ""), pal, fill=pal["green"], size=0.48)
        add_text(slide, x + 0.66, yy + 0.03, w - 0.72, 0.44, txt, size=15.0, color=pal["dark"])


def add_pipeline(slide, x, y, w, h, pal, variant="white"):
    add_panel(slide, x, y, w, h, "Methodology Pipeline", 3, pal, fill=(247, 252, 251), title_fill=pal["blue"] if variant == "minimal" else pal["teal"])
    steps = [
        ("Data collection", "OPSD hourly data", "database"),
        ("Preprocessing", "interpolate, align UTC", "audit"),
        ("Feature engineering", "lags, RES, TSO zones", "ai"),
        ("Modeling", "LR, SVM, RF, HGB, NN", "ai"),
        ("Evaluation", "F1, AUC, runtime", "chart"),
        ("Insights", "Pareto + guidance", "leaf"),
    ]
    top = y + 1.25
    step_gap = (h - 2.2) / (len(steps) - 1)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 0.92), Inches(top + 0.36), Inches(x + 0.92), Inches(top + step_gap * (len(steps) - 1) + 0.36))
    line.line.color.rgb = rgb((172, 195, 200))
    line.line.width = Pt(2)
    for i, (title, sub, icon) in enumerate(steps):
        yy = top + i * step_gap
        add_icon(slide, x + 0.48, yy, icon, pal, scale=1.05, color=pal["teal"] if i % 2 == 0 else pal["blue"])
        add_text(slide, x + 1.52, yy + 0.02, w - 1.75, 0.40, title, size=16.0, bold=True, color=pal["dark"])
        add_text(slide, x + 1.52, yy + 0.48, w - 1.75, 0.34, sub, size=11.8, color=pal["muted"])
        if i < len(steps) - 1:
            tri = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, Inches(x + 0.78), Inches(yy + step_gap * 0.52), Inches(0.28), Inches(0.42))
            tri.fill.solid()
            tri.fill.fore_color.rgb = rgb(pal["green"])
            tri.line.fill.background()


def add_dataset_cards(slide, x, y, w, pal):
    data = [
        ("OPSD Time Series", "~50,401 raw hourly rows", "load, wind, solar, TSO zones", "database"),
        ("Processed ML Table", "50,209 rows; 29 features", "lags, residual-load proxy", "chart"),
        ("Artifacts", "5 notebooks; 10 figures", "CSV tables + IEEE report", "doc"),
    ]
    card_w = (w - 0.5) / 3
    for i, (title, metric, detail, icon) in enumerate(data):
        xx = x + i * (card_w + 0.25)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(xx), Inches(y), Inches(card_w), Inches(2.42))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb((249, 253, 252))
        shape.line.color.rgb = rgb(pal["line"])
        add_icon(slide, xx + 0.25, y + 0.34, icon, pal, scale=0.86, color=pal["blue"] if i != 1 else pal["green"])
        add_text(slide, xx + 1.08, y + 0.24, card_w - 1.25, 0.50, title, size=15.5, bold=True, color=pal["dark"])
        add_text(slide, xx + 1.08, y + 0.83, card_w - 1.25, 0.46, metric, size=14.2, bold=True, color=pal["teal"])
        add_text(slide, xx + 1.08, y + 1.36, card_w - 1.25, 0.72, detail, size=12.0, color=pal["muted"])


def add_bar_chart(slide, x, y, w, h, categories, values, title, pal, color=None):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(title, values)
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = chart_shape.chart
    chart.has_legend = False
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(15)
    chart.chart_title.text_frame.paragraphs[0].font.bold = True
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.number_format = "0.0" if "F1" in title else "0"
    if "F1" in title:
        chart.value_axis.minimum_scale = 0.0
        chart.value_axis.maximum_scale = 1.0
    if color:
        chart.plots[0].series[0].format.fill.solid()
        chart.plots[0].series[0].format.fill.fore_color.rgb = rgb(color)
    return chart_shape


def add_column_chart(slide, x, y, w, h, categories, values, title, pal, color=None):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(title, values)
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = chart_shape.chart
    chart.has_legend = False
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(15)
    chart.chart_title.text_frame.paragraphs[0].font.bold = True
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.number_format = "0.0" if "F1" in title else "0"
    if "F1" in title:
        chart.value_axis.minimum_scale = 0.0
        chart.value_axis.maximum_scale = 1.0
    if color:
        chart.plots[0].series[0].format.fill.solid()
        chart.plots[0].series[0].format.fill.fore_color.rgb = rgb(color)
    return chart_shape


def add_top_header(slide, pal, variant):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(pal["bg"])
    header_color = pal["header"]
    header_h = 3.18 if variant != "dark" else 3.45
    header = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(header_h))
    header.fill.solid()
    header.fill.fore_color.rgb = rgb(header_color)
    header.line.fill.background()
    title_color = (255, 255, 255) if variant == "dark" else pal["dark"]
    subtitle_color = (200, 230, 218) if variant == "dark" else pal["muted"]
    add_text(
        slide,
        5.0,
        0.30,
        38.0,
        1.00,
        "Auditing Carbon-Aware AI Model Selection",
        size=56,
        bold=True,
        color=title_color,
        align=PP_ALIGN.CENTER,
        font="Aptos Display",
    )
    add_text(
        slide,
        5.0,
        1.34,
        38.0,
        0.58,
        "Accuracy, Sustainability, Transparency and Robustness in German Power-System Analytics",
        size=30,
        bold=True,
        color=title_color,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        5.0,
        2.14,
        38.0,
        0.42,
        "AI, Power & Responsibility - Governing Intelligent Systems for Sustainability",
        size=17,
        color=subtitle_color,
        align=PP_ALIGN.CENTER,
    )
    for i, (kind, xx) in enumerate([("ai", 1.2), ("leaf", 2.1), ("database", 44.9), ("audit", 45.8)]):
        add_icon(slide, xx, 0.82 if variant != "dark" else 0.98, kind, pal, scale=1.45, color=pal["green"] if i in (1, 3) else pal["blue"])


def add_footer(slide, pal, variant):
    if variant == "dark":
        fill = pal["header"]
        text_color = (212, 238, 230)
    else:
        fill = pal["teal"]
        text_color = (255, 255, 255)
    footer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(26.25), Inches(W), Inches(0.75))
    footer.fill.solid()
    footer.fill.fore_color.rgb = rgb(fill)
    footer.line.fill.background()
    add_text(slide, 1.1, 26.37, 44.8, 0.34, "Dataset: Open Power System Data Time Series (2020-10-06) | Reproducible notebooks, figures, tables, workflow, and IEEE/PES report", size=12.0, color=text_color, align=PP_ALIGN.CENTER)


def add_main_results(slide, x, y, w, h, pal, compact=False):
    rq1 = pd.read_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv")
    rq2 = pd.read_csv(TABLE_DIR / "rq2_feature_ablation_metrics.csv")
    rq5 = pd.read_csv(TABLE_DIR / "rq5_training_window_scenarios.csv")

    model_labels = ["HGB", "NN", "RF", "LR", "SVM"]
    f1_vals = [0.899, 0.897, 0.890, 0.728, 0.723]
    add_bar_chart(slide, x, y, w * 0.46, h * 0.43, model_labels, f1_vals, "F1 by model", pal, color=pal["teal"])

    feature_labels = ["Cal", "+Load", "+RES", "+TSO", "Full"]
    feature_vals = [0.813, 0.904, 0.889, 0.902, 0.899]
    add_column_chart(slide, x + w * 0.50, y, w * 0.48, h * 0.43, feature_labels, feature_vals, "Feature ablation F1", pal, color=pal["green"])

    carbon = rq1.set_index("model").loc[["Linear SVM", "Logistic Regression", "Histogram Gradient Boosting", "Random Forest", "Neural Network"], "training_carbon_mgco2e_proxy"].round(2)
    add_bar_chart(slide, x, y + h * 0.50, w * 0.46, h * 0.42, ["SVM", "LR", "HGB", "RF", "NN"], carbon.tolist(), "Carbon proxy (mg CO2e)", pal, color=pal["blue"])

    green_reduction = rq5[rq5["scenario"].str.startswith("Green")]["reduction_vs_random_percent"].iloc[0]
    add_metric(slide, x + w * 0.55, y + h * 0.55, w * 0.16, "green-window reduction", f"{green_reduction:.1f}%", pal, pal["green"])
    add_metric(slide, x + w * 0.72, y + h * 0.55, w * 0.13, "best F1", "0.899", pal, pal["teal"])
    add_metric(slide, x + w * 0.85, y + h * 0.55, w * 0.12, "ROC-AUC", "0.993", pal, pal["blue"])
    add_text(slide, x + w * 0.52, y + h * 0.78, w * 0.44, 0.92, "HGB is the balanced recommendation; Linear SVM is the lowest-compute fallback.", size=14.5, bold=True, color=pal["dark"], align=PP_ALIGN.CENTER)


def add_poster(slide, pal, variant="white"):
    add_top_header(slide, pal, variant)
    top_y = 3.62 if variant != "dark" else 3.86

    add_panel(slide, 1.0, top_y, 22.2, 4.75, "Motivation", 1, pal, fill=pal["panel"], title_fill=pal["green"] if variant != "dark" else pal["teal"])
    add_bullets(
        slide,
        1.42,
        top_y + 0.85,
        20.8,
        [
            "Accuracy-only ML model choice hides compute and carbon costs.",
            "German grid analytics needs robust, season-aware evidence.",
            "Transparent reporting should include F1, AUC, runtime and size.",
            "Carbon-aware scheduling can reduce avoidable training impact.",
        ],
        pal,
            size=16.0,
            gap=0.83,
    )

    add_panel(slide, 24.0, top_y, 23.0, 4.75, "Datasets", 2, pal, fill=pal["panel"], title_fill=pal["blue"])
    add_dataset_cards(slide, 24.45, top_y + 0.95, 22.1, pal)

    left_y = top_y + 5.25
    add_pipeline(slide, 1.0, left_y, 8.7, 16.25, pal, variant)

    main_x = 10.25
    add_panel(slide, main_x, left_y, 12.9, 7.4, "Research Questions", 4, pal, fill=pal["panel"], title_fill=pal["teal"])
    add_rq_list(slide, main_x + 0.48, left_y + 0.88, 11.8, pal)

    add_panel(slide, 23.8, left_y, 23.2, 10.5, "Key Results (Highlights)", 5, pal, fill=pal["panel"], title_fill=pal["blue"])
    add_main_results(slide, 24.35, left_y + 0.83, 22.1, 8.85, pal)

    add_panel(slide, main_x, left_y + 7.9, 12.9, 5.35, "Key Findings", 6, pal, fill=pal["panel"], title_fill=pal["green"])
    findings = [
        "HGB: best F1 and Pareto-efficient.",
        "Load history gives the largest feature gain.",
        "Summer is the hardest seasonal regime.",
        "Use low-compute baselines for auditability.",
    ]
    add_bullets(slide, main_x + 0.48, left_y + 8.75, 11.8, findings, pal, icon_kind="audit", size=14.0, gap=0.75)

    add_panel(slide, 23.8, left_y + 11.0, 23.2, 5.25, "Conclusion", 7, pal, fill=pal["panel"], title_fill=pal["teal"])
    add_text(slide, 24.35, left_y + 11.82, 22.0, 1.28, "Responsible model selection should be Pareto-based: accuracy, carbon proxy, runtime, model size, seasonal robustness and auditability are evaluated together.", size=18.0, bold=True, color=pal["dark"], align=PP_ALIGN.CENTER)
    add_text(slide, 24.55, left_y + 13.55, 21.5, 1.02, "Recommended balanced model: Histogram Gradient Boosting. Recommended low-compute fallback: Linear SVM. Next step: replace proxy energy with measured hardware power.", size=13.5, color=pal["muted"], align=PP_ALIGN.CENTER)

    add_footer(slide, pal, variant)


def create_poster_deck() -> None:
    POSTER_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for variant in ["white", "dark", "minimal"]:
        slide = prs.slides.add_slide(blank)
        add_poster(slide, PALETTES[variant], variant=variant)
    prs.save(OUT_PPTX)


def draw_workflow_icon(ax, kind, x, y, scale=1.0, color="#0b8793"):
    if kind == "database":
        ax.add_patch(Ellipse((x, y + 0.07 * scale), 0.36 * scale, 0.11 * scale, fc="#d7eef4", ec=color, lw=1.6))
        ax.add_patch(Rectangle((x - 0.18 * scale, y - 0.12 * scale), 0.36 * scale, 0.19 * scale, fc="#edf8fb", ec=color, lw=1.6))
        ax.add_patch(Ellipse((x, y - 0.12 * scale), 0.36 * scale, 0.11 * scale, fc="#edf8fb", ec=color, lw=1.6))
    elif kind == "api":
        ax.add_patch(Rectangle((x - 0.17 * scale, y - 0.14 * scale), 0.34 * scale, 0.28 * scale, fc="#edf8fb", ec=color, lw=1.6))
        ax.text(x, y, "API", ha="center", va="center", fontsize=8 * scale, color=color, fontweight="bold")
    elif kind == "features":
        for i, lab in enumerate(["F1", "F2", "F3", "Fn"]):
            ax.add_patch(Rectangle((x - 0.16 * scale, y + (0.19 - i * 0.12) * scale), 0.32 * scale, 0.08 * scale, fc="#e8f7ef", ec=color, lw=1.1))
            ax.text(x, y + (0.23 - i * 0.12) * scale, lab, ha="center", va="center", fontsize=5.5 * scale, color="#17324d")
    elif kind == "globe":
        ax.add_patch(Circle((x, y), 0.18 * scale, fc="#e8f7ef", ec=color, lw=1.6))
        ax.plot([x - 0.18 * scale, x + 0.18 * scale], [y, y], color=color, lw=0.9)
        ax.plot([x, x], [y - 0.18 * scale, y + 0.18 * scale], color=color, lw=0.9)
        ax.add_patch(Ellipse((x, y), 0.16 * scale, 0.36 * scale, fc="none", ec=color, lw=0.9))
    elif kind == "model":
        ax.add_patch(Rectangle((x - 0.18 * scale, y - 0.15 * scale), 0.36 * scale, 0.30 * scale, fc="#e8f7ef", ec=color, lw=1.6))
        for off in [-0.07, 0.07]:
            ax.plot([x - 0.14 * scale, x + 0.14 * scale], [y + off * scale, y + off * scale], color=color, lw=1)
            ax.plot([x + off * scale, x + off * scale], [y - 0.12 * scale, y + 0.12 * scale], color=color, lw=1)
    elif kind == "chart":
        ax.plot([x - 0.18 * scale, x - 0.18 * scale, x + 0.18 * scale], [y + 0.15 * scale, y - 0.15 * scale, y - 0.15 * scale], color=color, lw=1.4)
        for dx, dy, c in [(-0.08, -0.05, "#2a9d8f"), (0.03, 0.04, "#3064a2"), (0.13, 0.11, "#46ae5e")]:
            ax.add_patch(Circle((x + dx * scale, y + dy * scale), 0.035 * scale, fc=c, ec="#17324d", lw=0.7))
    elif kind == "leaf":
        ax.add_patch(Ellipse((x, y), 0.33 * scale, 0.20 * scale, angle=35, fc="#78c68b", ec=color, lw=1.6))
        ax.plot([x - 0.10 * scale, x + 0.10 * scale], [y - 0.06 * scale, y + 0.06 * scale], color="white", lw=1.5)


def create_structured_workflow() -> None:
    steps = [
        ("Data source", "OPSD time-series CSV\nmetadata and IEEE/PES prompt", "database"),
        ("Data ingestion", "data.csv -> reproducible\nPython project workspace", "api"),
        ("Feature extraction", "F1-Fn: calendar, lags,\nRES, TSO zones, forecast", "features"),
        ("Segmentation", "feature families and\nseasonal operating regimes", "globe"),
        ("Complete Dataset", "preprocessed_content.csv\n50,209 rows; 29 features", "database"),
        ("AI Framework", "normalization, feature selection,\nmodel training and prediction", "model"),
        ("Output results", "high-load risk, model ranking,\ncarbon proxy and figures", "chart"),
        ("Recommendations", "Pareto selection, green training\nwindows, audit documentation", "leaf"),
    ]
    fig, ax = plt.subplots(figsize=(18.2, 7.8), dpi=220)
    ax.set_xlim(0, 18.2)
    ax.set_ylim(0, 7.8)
    ax.axis("off")
    ax.text(9.1, 7.35, "Structured AI Workflow for Carbon-Aware Sustainability Auditing", ha="center", va="center", fontsize=17.5, fontweight="bold", color="#102a43")
    ax.text(9.1, 7.02, "Germany-aligned power-system data -> carbon-aware ML evaluation -> IEEE/PES-ready evidence and recommendations", ha="center", va="center", fontsize=9.5, color="#52616b")

    positions = [(1.0, 4.68), (3.05, 4.68), (5.1, 4.68), (7.15, 4.68), (9.2, 4.68), (11.95, 4.68), (14.55, 4.68), (17.0, 4.68)]
    box_w, box_h = 1.62, 1.12
    colors = ["#0b8793", "#3064a2", "#2a9d8f", "#46ae5e"]

    # Large AI framework area.
    framework = FancyBboxPatch((10.35, 2.05), 3.25, 4.45, boxstyle="round,pad=0.06,rounding_size=0.06", fc="#f7fcfb", ec="#52616b", lw=1.7)
    ax.add_patch(framework)
    ax.text(11.98, 6.28, "AI Framework", ha="center", va="center", fontsize=10.0, fontweight="bold", color="#52616b")
    subblocks = [
        ("6.1", "Preprocessing", "normalization\\nscaling"),
        ("6.2", "Feature selection", "MI, PCA,\\npermutation"),
        ("6.3", "Model training", "LR, SVM, RF,\\nHGB, NN"),
        ("6.4", "Prediction", "next-day\\nhigh-load risk"),
    ]
    for i, (num, title, sub) in enumerate(subblocks):
        xx = 10.58 + (i % 2) * 1.46
        yy = 3.70 - (i // 2) * 1.10
        ax.add_patch(FancyBboxPatch((xx, yy - 0.88), 1.16, 0.78, boxstyle="round,pad=0.04,rounding_size=0.04", fc="white", ec="#9fb7bd", lw=1.0))
        ax.add_patch(Circle((xx + 0.09, yy - 0.08), 0.16, fc="#0b8793", ec="#17324d", lw=0.9))
        ax.text(xx + 0.09, yy - 0.08, num.split(".")[1], ha="center", va="center", fontsize=7, color="white", fontweight="bold")
        ax.text(xx + 0.62, yy - 0.28, title, ha="center", va="center", fontsize=7.2, fontweight="bold", color="#102a43")
        ax.text(xx + 0.62, yy - 0.58, sub, ha="center", va="center", fontsize=5.8, color="#52616b")

    for idx, ((title, subtitle, icon), (x, y)) in enumerate(zip(steps, positions), start=1):
        if idx == 6:
            box_fc = "#e8f7ef"
        else:
            box_fc = "#ffffff"
        ax.add_patch(FancyBboxPatch((x - box_w / 2, y - box_h / 2), box_w, box_h, boxstyle="round,pad=0.05,rounding_size=0.06", fc=box_fc, ec="#17324d", lw=1.35))
        c = colors[(idx - 1) % len(colors)]
        ax.add_patch(Circle((x - box_w / 2 + 0.05, y + box_h / 2 + 0.17), 0.23, fc=c, ec="#17324d", lw=1.1))
        ax.text(x - box_w / 2 + 0.05, y + box_h / 2 + 0.17, str(idx), ha="center", va="center", fontsize=10.5, color="white", fontweight="bold")
        draw_workflow_icon(ax, icon, x, y + 0.18, 1.25, color=c)
        ax.text(x, y - 0.13, title, ha="center", va="center", fontsize=7.7, fontweight="bold", color="#102a43")
        ax.text(x, y - 0.43, subtitle, ha="center", va="center", fontsize=5.6, color="#52616b", linespacing=1.08)

    def arrow(a, b, rad=0):
        ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=14, lw=1.65, color="#17324d", connectionstyle=f"arc3,rad={rad}"))

    for i in range(7):
        x1, y1 = positions[i]
        x2, y2 = positions[i + 1]
        arrow((x1 + box_w / 2 + 0.08, y1), (x2 - box_w / 2 - 0.08, y2))
    ax.text(9.1, 0.92, "Reproducible outputs: raw/processed data, five notebooks, five CSV result tables, publication figures, vector icons, editable posters, documentation, and IEEE/PES technical report.", ha="center", va="center", fontsize=7.8, color="#52616b")

    FIG_DIR.mkdir(parents=True, exist_ok=True)
    for ext in ("png", "pdf"):
        fig.savefig(FIG_DIR / f"workflow_structured_ai_esg_style.{ext}", bbox_inches="tight", facecolor="white", dpi=300)
    plt.close(fig)


def main() -> None:
    create_structured_workflow()
    create_poster_deck()
    print(OUT_PPTX)
    print(FIG_DIR / "workflow_structured_ai_esg_style.png")


if __name__ == "__main__":
    main()
