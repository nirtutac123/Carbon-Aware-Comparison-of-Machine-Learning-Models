from __future__ import annotations

import json
import math
import pickle
import textwrap
import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from sklearn.ensemble import HistGradientBoostingClassifier, RandomForestClassifier
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    f1_score,
    precision_score,
    recall_score,
    roc_auc_score,
)
from sklearn.neural_network import MLPClassifier
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import LinearSVC


ROOT = Path(__file__).resolve().parents[1]
RAW_DATA = ROOT / "data.csv"
PREPROCESSED = ROOT / "preprocessed_content.csv"
FIG_DIR = ROOT / "outputs" / "figures"
TABLE_DIR = ROOT / "outputs" / "tables"
NOTEBOOK_DIR = ROOT / "notebooks"
ICON_DIR = ROOT / "icons"
POSTER_DIR = ROOT / "posters"

RANDOM_STATE = 42
TRAIN_CPU_WATTS = 45.0
INFER_CPU_WATTS = 18.0

CALENDAR_FEATURES = [
    "hour_sin",
    "hour_cos",
    "dow_sin",
    "dow_cos",
    "doy_sin",
    "doy_cos",
    "is_weekend",
]
LOAD_HISTORY_FEATURES = [
    "load_mw",
    "load_lag_1h",
    "load_lag_24h",
    "load_lag_168h",
    "load_roll_24h_mean",
    "load_roll_168h_mean",
    "load_roll_24h_std",
]
RENEWABLE_FEATURES = [
    "solar_generation_mw",
    "wind_generation_mw",
    "wind_onshore_generation_mw",
    "wind_offshore_generation_mw",
    "renewable_generation_mw",
    "renewable_share",
    "renewable_share_lag_24h",
    "residual_load_mw",
    "residual_load_share",
]
ZONE_FEATURES = [
    "zone_50hertz_load_mw",
    "zone_amprion_load_mw",
    "zone_tennet_load_mw",
    "zone_transnetbw_load_mw",
    "zone_load_imbalance_mw",
]
FORECAST_FEATURES = ["load_forecast_mw"]
ALL_FEATURES = (
    CALENDAR_FEATURES
    + LOAD_HISTORY_FEATURES
    + RENEWABLE_FEATURES
    + ZONE_FEATURES
    + FORECAST_FEATURES
)


RESEARCH_QUESTIONS = [
    {
        "id": "RQ1",
        "question": (
            "Which model family gives the strongest Pareto trade-off between "
            "predictive accuracy and estimated computational carbon for German "
            "next-day high-load-event classification?"
        ),
        "figure": "rq1_accuracy_carbon_pareto.pdf",
        "table": "rq1_model_tradeoff_metrics.csv",
    },
    {
        "id": "RQ2",
        "question": (
            "How much do load history, renewable generation, TSO-zone state, "
            "and day-ahead forecast features improve performance relative to a "
            "calendar-only baseline?"
        ),
        "figure": "rq2_feature_ablation.pdf",
        "table": "rq2_feature_ablation_metrics.csv",
    },
    {
        "id": "RQ3",
        "question": (
            "Are the model rankings stable across German winter, spring, "
            "summer, and autumn operating regimes?"
        ),
        "figure": "rq3_seasonal_f1_heatmap.pdf",
        "table": "rq3_seasonal_generalization.csv",
    },
    {
        "id": "RQ4",
        "question": (
            "Can an interpretable low-complexity model explain the dominant "
            "drivers of high-load risk without a large loss in predictive utility?"
        ),
        "figure": "rq4_interpretability_efficiency.pdf",
        "table": "rq4_top_interpretable_features.csv",
    },
    {
        "id": "RQ5",
        "question": (
            "How much carbon-proxy reduction is available if model training is "
            "scheduled into low-residual-load windows rather than arbitrary or "
            "high-residual-load hours?"
        ),
        "figure": "rq5_carbon_aware_training_windows.pdf",
        "table": "rq5_training_window_scenarios.csv",
    },
]


def ensure_dirs() -> None:
    for path in [FIG_DIR, TABLE_DIR, NOTEBOOK_DIR, ICON_DIR, POSTER_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def season_name(month: int) -> str:
    if month in (12, 1, 2):
        return "Winter"
    if month in (3, 4, 5):
        return "Spring"
    if month in (6, 7, 8):
        return "Summer"
    return "Autumn"


def save_figure(fig: plt.Figure, stem: str) -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    for ext in ("pdf", "png"):
        fig.savefig(
            FIG_DIR / f"{stem}.{ext}",
            bbox_inches="tight",
            dpi=300,
            facecolor="white",
        )
    plt.close(fig)


def load_or_create_preprocessed(force: bool = False) -> pd.DataFrame:
    ensure_dirs()
    if PREPROCESSED.exists() and not force:
        return pd.read_csv(PREPROCESSED, parse_dates=["timestamp_utc"])

    if not RAW_DATA.exists():
        raise FileNotFoundError(
            f"Expected raw OPSD file at {RAW_DATA}. Download it as data.csv first."
        )

    columns = [
        "utc_timestamp",
        "cet_cest_timestamp",
        "DE_load_actual_entsoe_transparency",
        "DE_load_forecast_entsoe_transparency",
        "DE_solar_generation_actual",
        "DE_wind_generation_actual",
        "DE_wind_offshore_generation_actual",
        "DE_wind_onshore_generation_actual",
        "DE_50hertz_load_actual_entsoe_transparency",
        "DE_amprion_load_actual_entsoe_transparency",
        "DE_tennet_load_actual_entsoe_transparency",
        "DE_transnetbw_load_actual_entsoe_transparency",
    ]
    df = pd.read_csv(RAW_DATA, usecols=columns)
    df = df.rename(
        columns={
            "utc_timestamp": "timestamp_utc",
            "DE_load_actual_entsoe_transparency": "load_mw",
            "DE_load_forecast_entsoe_transparency": "load_forecast_mw",
            "DE_solar_generation_actual": "solar_generation_mw",
            "DE_wind_generation_actual": "wind_generation_mw",
            "DE_wind_offshore_generation_actual": "wind_offshore_generation_mw",
            "DE_wind_onshore_generation_actual": "wind_onshore_generation_mw",
            "DE_50hertz_load_actual_entsoe_transparency": "zone_50hertz_load_mw",
            "DE_amprion_load_actual_entsoe_transparency": "zone_amprion_load_mw",
            "DE_tennet_load_actual_entsoe_transparency": "zone_tennet_load_mw",
            "DE_transnetbw_load_actual_entsoe_transparency": "zone_transnetbw_load_mw",
        }
    )
    df["timestamp_utc"] = pd.to_datetime(df["timestamp_utc"], utc=True)
    df = df.sort_values("timestamp_utc").reset_index(drop=True)

    numeric_cols = [c for c in df.columns if c not in ("timestamp_utc", "cet_cest_timestamp")]
    df[numeric_cols] = df[numeric_cols].interpolate(limit_direction="both")
    df[numeric_cols] = df[numeric_cols].ffill().bfill()

    dt = df["timestamp_utc"].dt
    df["year"] = dt.year
    df["month"] = dt.month
    df["hour"] = dt.hour
    df["dayofweek"] = dt.dayofweek
    df["dayofyear"] = dt.dayofyear
    df["is_weekend"] = (df["dayofweek"] >= 5).astype(int)
    df["season"] = df["month"].map(season_name)
    df["hour_sin"] = np.sin(2 * np.pi * df["hour"] / 24)
    df["hour_cos"] = np.cos(2 * np.pi * df["hour"] / 24)
    df["dow_sin"] = np.sin(2 * np.pi * df["dayofweek"] / 7)
    df["dow_cos"] = np.cos(2 * np.pi * df["dayofweek"] / 7)
    df["doy_sin"] = np.sin(2 * np.pi * df["dayofyear"] / 366)
    df["doy_cos"] = np.cos(2 * np.pi * df["dayofyear"] / 366)

    df["renewable_generation_mw"] = (
        df["solar_generation_mw"] + df["wind_generation_mw"]
    )
    df["renewable_share"] = (
        df["renewable_generation_mw"] / df["load_mw"].replace(0, np.nan)
    ).clip(lower=0, upper=1.5)
    df["residual_load_mw"] = (df["load_mw"] - df["renewable_generation_mw"]).clip(
        lower=0
    )
    df["residual_load_share"] = (
        df["residual_load_mw"] / df["load_mw"].replace(0, np.nan)
    ).clip(lower=0, upper=1.2)
    df["carbon_intensity_proxy_g_per_kwh"] = 150 + 500 * df[
        "residual_load_share"
    ].clip(0, 1)

    for lag in (1, 24, 168):
        df[f"load_lag_{lag}h"] = df["load_mw"].shift(lag)
    df["load_roll_24h_mean"] = df["load_mw"].rolling(24, min_periods=24).mean()
    df["load_roll_168h_mean"] = df["load_mw"].rolling(168, min_periods=168).mean()
    df["load_roll_24h_std"] = df["load_mw"].rolling(24, min_periods=24).std()
    df["renewable_share_lag_24h"] = df["renewable_share"].shift(24)

    zone_cols = [
        "zone_50hertz_load_mw",
        "zone_amprion_load_mw",
        "zone_tennet_load_mw",
        "zone_transnetbw_load_mw",
    ]
    df["zone_load_imbalance_mw"] = df[zone_cols].max(axis=1) - df[zone_cols].min(
        axis=1
    )

    df["target_load_24h_mw"] = df["load_mw"].shift(-24)
    threshold = df["target_load_24h_mw"].quantile(0.80)
    df["target_high_load_next_24h"] = (
        df["target_load_24h_mw"] >= threshold
    ).astype(int)

    needed = ALL_FEATURES + [
        "target_high_load_next_24h",
        "target_load_24h_mw",
        "carbon_intensity_proxy_g_per_kwh",
    ]
    df = df.replace([np.inf, -np.inf], np.nan).dropna(subset=needed).reset_index(
        drop=True
    )

    out_cols = [
        "timestamp_utc",
        "cet_cest_timestamp",
        "year",
        "month",
        "season",
        "hour",
        "dayofweek",
        "dayofyear",
        *ALL_FEATURES,
        "target_load_24h_mw",
        "target_high_load_next_24h",
        "carbon_intensity_proxy_g_per_kwh",
    ]
    df[out_cols].to_csv(PREPROCESSED, index=False)

    metadata = {
        "raw_dataset": "Open Power System Data time_series_60min_singleindex.csv",
        "raw_file": "data.csv",
        "processed_file": "preprocessed_content.csv",
        "target": "target_high_load_next_24h",
        "target_definition": (
            "1 if German load 24 hours ahead is at or above the 80th percentile "
            f"of observed 24-hour-ahead load ({threshold:.1f} MW)."
        ),
        "date_range_utc": [
            str(df["timestamp_utc"].min()),
            str(df["timestamp_utc"].max()),
        ],
        "n_rows": int(len(df)),
        "n_features": len(ALL_FEATURES),
        "carbon_proxy_note": (
            "The carbon variable is a residual-load proxy, not a measured "
            "German grid-emission inventory. It scales from 150 to 650 gCO2e/kWh "
            "as the load share unmet by observed wind and solar increases."
        ),
    }
    (TABLE_DIR / "dataset_metadata.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )
    return df[out_cols]


def train_test_split_time(df: pd.DataFrame, features: list[str]):
    train_mask = df["timestamp_utc"] < pd.Timestamp("2019-01-01", tz="UTC")
    test_mask = ~train_mask
    X_train = df.loc[train_mask, features].astype(float)
    X_test = df.loc[test_mask, features].astype(float)
    y_train = df.loc[train_mask, "target_high_load_next_24h"].astype(int)
    y_test = df.loc[test_mask, "target_high_load_next_24h"].astype(int)
    return X_train, X_test, y_train, y_test, df.loc[test_mask].copy()


def build_models() -> dict[str, object]:
    return {
        "Logistic Regression": make_pipeline(
            StandardScaler(),
            LogisticRegression(
                max_iter=2500,
                class_weight="balanced",
                solver="lbfgs",
                random_state=RANDOM_STATE,
            ),
        ),
        "Linear SVM": make_pipeline(
            StandardScaler(),
            LinearSVC(
                max_iter=6000,
                class_weight="balanced",
                random_state=RANDOM_STATE,
                dual="auto",
            ),
        ),
        "Random Forest": RandomForestClassifier(
            n_estimators=120,
            max_depth=14,
            min_samples_leaf=20,
            class_weight="balanced_subsample",
            random_state=RANDOM_STATE,
            n_jobs=-1,
        ),
        "Histogram Gradient Boosting": HistGradientBoostingClassifier(
            max_iter=140,
            learning_rate=0.055,
            max_leaf_nodes=31,
            l2_regularization=0.05,
            random_state=RANDOM_STATE,
        ),
        "Neural Network": make_pipeline(
            StandardScaler(),
            MLPClassifier(
                hidden_layer_sizes=(32,),
                alpha=0.003,
                early_stopping=True,
                validation_fraction=0.12,
                max_iter=180,
                random_state=RANDOM_STATE,
            ),
        ),
    }


def predict_score(model, X: pd.DataFrame) -> np.ndarray:
    if hasattr(model, "predict_proba"):
        return model.predict_proba(X)[:, 1]
    if hasattr(model, "decision_function"):
        score = model.decision_function(X)
        return 1 / (1 + np.exp(-np.clip(score, -30, 30)))
    return model.predict(X)


def model_size_kb(model) -> float:
    return len(pickle.dumps(model)) / 1024


def carbon_metrics(fit_time_s: float, pred_time_s: float, mean_proxy: float) -> tuple[float, float]:
    energy_wh = fit_time_s * TRAIN_CPU_WATTS / 3600 + pred_time_s * INFER_CPU_WATTS / 3600
    carbon_mg = energy_wh * mean_proxy
    return energy_wh, carbon_mg


def evaluate_models(
    df: pd.DataFrame,
    features: list[str],
    model_dict: dict[str, object] | None = None,
) -> tuple[pd.DataFrame, dict[str, object], pd.DataFrame, pd.Series]:
    if model_dict is None:
        model_dict = build_models()
    X_train, X_test, y_train, y_test, test_df = train_test_split_time(df, features)
    rows = []
    fitted = {}
    mean_proxy = float(df.loc[df["timestamp_utc"] < pd.Timestamp("2019-01-01", tz="UTC"), "carbon_intensity_proxy_g_per_kwh"].mean())

    for name, model in model_dict.items():
        t0 = time.perf_counter()
        model.fit(X_train, y_train)
        fit_time = time.perf_counter() - t0

        t1 = time.perf_counter()
        y_pred = model.predict(X_test)
        y_score = predict_score(model, X_test)
        pred_time = time.perf_counter() - t1

        energy_wh, carbon_mg = carbon_metrics(fit_time, pred_time, mean_proxy)
        rows.append(
            {
                "model": name,
                "n_train": int(len(X_train)),
                "n_test": int(len(X_test)),
                "accuracy": accuracy_score(y_test, y_pred),
                "balanced_accuracy": balanced_accuracy_score(y_test, y_pred),
                "precision": precision_score(y_test, y_pred, zero_division=0),
                "recall": recall_score(y_test, y_pred, zero_division=0),
                "f1_score": f1_score(y_test, y_pred, zero_division=0),
                "roc_auc": roc_auc_score(y_test, y_score),
                "fit_time_s": fit_time,
                "predict_time_s": pred_time,
                "model_size_kb": model_size_kb(model),
                "energy_wh_proxy": energy_wh,
                "training_carbon_mgco2e_proxy": carbon_mg,
                "mean_training_carbon_proxy_g_per_kwh": mean_proxy,
            }
        )
        fitted[name] = model
    metrics = pd.DataFrame(rows).sort_values("f1_score", ascending=False)
    return metrics, fitted, test_df, y_test


def mark_pareto(metrics: pd.DataFrame) -> pd.DataFrame:
    out = metrics.copy()
    efficient = []
    for _, row in out.iterrows():
        dominated = (
            (out["f1_score"] >= row["f1_score"])
            & (
                out["training_carbon_mgco2e_proxy"]
                <= row["training_carbon_mgco2e_proxy"]
            )
            & (
                (out["f1_score"] > row["f1_score"])
                | (
                    out["training_carbon_mgco2e_proxy"]
                    < row["training_carbon_mgco2e_proxy"]
                )
            )
        ).any()
        efficient.append(not dominated)
    out["pareto_efficient"] = efficient
    out["carbon_adjusted_score"] = (
        out["f1_score"]
        - 0.025 * np.log10(1 + out["training_carbon_mgco2e_proxy"])
        - 0.010 * np.log10(1 + out["model_size_kb"])
    )
    return out


def run_rq1(df: pd.DataFrame, return_artifacts: bool = False):
    metrics, fitted, test_df, y_test = evaluate_models(df, ALL_FEATURES)
    metrics = mark_pareto(metrics)
    metrics.to_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv", index=False)

    fig, ax = plt.subplots(figsize=(8.4, 5.0))
    colors = ["#1b9aaa", "#ef476f", "#06d6a0", "#f4a261", "#6954a3"]
    for idx, row in metrics.reset_index(drop=True).iterrows():
        ax.scatter(
            row["training_carbon_mgco2e_proxy"],
            row["f1_score"],
            s=90 + 0.02 * row["model_size_kb"],
            color=colors[idx % len(colors)],
            edgecolor="black",
            linewidth=0.8,
            zorder=3,
        )
        label = row["model"].replace("Histogram ", "Hist. ")
        ax.annotate(
            label,
            (row["training_carbon_mgco2e_proxy"], row["f1_score"]),
            textcoords="offset points",
            xytext=(7, 7),
            fontsize=8.5,
        )
    frontier = metrics[metrics["pareto_efficient"]].sort_values(
        "training_carbon_mgco2e_proxy"
    )
    if len(frontier) > 1:
        ax.plot(
            frontier["training_carbon_mgco2e_proxy"],
            frontier["f1_score"],
            color="#2b2d42",
            linewidth=1.6,
            linestyle="--",
            label="Pareto frontier",
        )
    ax.set_xscale("log")
    ax.set_xlabel("Training carbon proxy (mg CO2e, log scale)")
    ax.set_ylabel("F1-score on 2019-2020 test period")
    ax.set_title("RQ1: Accuracy-Carbon Trade-off for German Next-Day High-Load Classification")
    ax.grid(True, which="both", alpha=0.24)
    ax.legend(frameon=False, loc="lower right")
    save_figure(fig, "rq1_accuracy_carbon_pareto")

    if return_artifacts:
        return metrics, fitted, test_df, y_test
    return metrics


def run_rq2(df: pd.DataFrame) -> pd.DataFrame:
    feature_sets = [
        ("Calendar only", CALENDAR_FEATURES),
        ("Calendar + national load history", CALENDAR_FEATURES + LOAD_HISTORY_FEATURES),
        (
            "Add renewable operating state",
            CALENDAR_FEATURES + LOAD_HISTORY_FEATURES + RENEWABLE_FEATURES,
        ),
        (
            "Add TSO-zone state",
            CALENDAR_FEATURES
            + LOAD_HISTORY_FEATURES
            + RENEWABLE_FEATURES
            + ZONE_FEATURES,
        ),
        ("Full feature set incl. day-ahead forecast", ALL_FEATURES),
    ]
    rows = []
    model_name = "Histogram Gradient Boosting"
    for label, features in feature_sets:
        metrics, _, _, _ = evaluate_models(
            df, features, {model_name: build_models()[model_name]}
        )
        row = metrics.iloc[0].to_dict()
        row["feature_set"] = label
        row["n_features"] = len(features)
        rows.append(row)
    result = pd.DataFrame(rows)
    result = result[
        [
            "feature_set",
            "n_features",
            "accuracy",
            "balanced_accuracy",
            "precision",
            "recall",
            "f1_score",
            "roc_auc",
            "fit_time_s",
            "energy_wh_proxy",
            "training_carbon_mgco2e_proxy",
            "model_size_kb",
        ]
    ]
    result.to_csv(TABLE_DIR / "rq2_feature_ablation_metrics.csv", index=False)

    fig, ax1 = plt.subplots(figsize=(9.4, 4.9))
    x = np.arange(len(result))
    ax1.bar(
        x,
        result["f1_score"],
        color="#2a9d8f",
        width=0.58,
        label="F1-score",
    )
    ax1.set_ylim(0, max(1.0, result["f1_score"].max() * 1.12))
    ax1.set_ylabel("F1-score")
    ax1.set_xticks(x)
    ax1.set_xticklabels(result["feature_set"], rotation=18, ha="right")
    ax1.grid(axis="y", alpha=0.25)

    ax2 = ax1.twinx()
    ax2.plot(
        x,
        result["training_carbon_mgco2e_proxy"],
        color="#e76f51",
        marker="o",
        linewidth=2,
        label="Carbon proxy",
    )
    ax2.set_ylabel("Training carbon proxy (mg CO2e)")
    fig.suptitle("RQ2: Feature Value Versus Computational Cost")
    fig.tight_layout()
    save_figure(fig, "rq2_feature_ablation")
    return result


def run_rq3(
    df: pd.DataFrame,
    fitted_models: dict[str, object] | None = None,
    test_df: pd.DataFrame | None = None,
    y_test: pd.Series | None = None,
) -> pd.DataFrame:
    if fitted_models is None or test_df is None or y_test is None:
        _, fitted_models, test_df, y_test = evaluate_models(df, ALL_FEATURES)
    X_test = test_df[ALL_FEATURES].astype(float)
    rows = []
    for model_name, model in fitted_models.items():
        y_pred = pd.Series(model.predict(X_test), index=test_df.index)
        for season in ["Winter", "Spring", "Summer", "Autumn"]:
            mask = test_df["season"] == season
            if mask.sum() == 0:
                continue
            y_true_s = test_df.loc[mask, "target_high_load_next_24h"].astype(int)
            y_pred_s = y_pred.loc[mask]
            rows.append(
                {
                    "model": model_name,
                    "season": season,
                    "n_test": int(mask.sum()),
                    "event_rate": y_true_s.mean(),
                    "f1_score": f1_score(y_true_s, y_pred_s, zero_division=0),
                    "balanced_accuracy": balanced_accuracy_score(y_true_s, y_pred_s),
                    "precision": precision_score(y_true_s, y_pred_s, zero_division=0),
                    "recall": recall_score(y_true_s, y_pred_s, zero_division=0),
                }
            )
    seasonal = pd.DataFrame(rows)
    seasonal.to_csv(TABLE_DIR / "rq3_seasonal_generalization.csv", index=False)

    pivot = seasonal.pivot(index="model", columns="season", values="f1_score")
    pivot = pivot[["Winter", "Spring", "Summer", "Autumn"]]
    fig, ax = plt.subplots(figsize=(8.7, 4.2))
    im = ax.imshow(pivot.values, cmap="viridis", vmin=0, vmax=1)
    ax.set_xticks(np.arange(pivot.shape[1]))
    ax.set_xticklabels(pivot.columns)
    ax.set_yticks(np.arange(pivot.shape[0]))
    ax.set_yticklabels([m.replace("Histogram ", "Hist. ") for m in pivot.index])
    for i in range(pivot.shape[0]):
        for j in range(pivot.shape[1]):
            val = pivot.values[i, j]
            ax.text(
                j,
                i,
                f"{val:.2f}",
                ha="center",
                va="center",
                color="white" if val < 0.72 else "black",
                fontsize=9,
            )
    cbar = fig.colorbar(im, ax=ax, fraction=0.035, pad=0.03)
    cbar.set_label("F1-score")
    ax.set_title("RQ3: Seasonal Generalization on German Operating Regimes")
    fig.tight_layout()
    save_figure(fig, "rq3_seasonal_f1_heatmap")
    return seasonal


def feature_label(name: str) -> str:
    replacements = {
        "load_mw": "Current German load",
        "load_forecast_mw": "Day-ahead load forecast",
        "load_lag_1h": "Load lag, 1 h",
        "load_lag_24h": "Load lag, 24 h",
        "load_lag_168h": "Load lag, 168 h",
        "load_roll_24h_mean": "Load rolling mean, 24 h",
        "load_roll_168h_mean": "Load rolling mean, 168 h",
        "load_roll_24h_std": "Load volatility, 24 h",
        "renewable_generation_mw": "Observed wind + solar",
        "renewable_share": "Wind/solar share of load",
        "residual_load_mw": "Residual load",
        "residual_load_share": "Residual-load share",
        "zone_load_imbalance_mw": "TSO load imbalance",
        "zone_50hertz_load_mw": "50Hertz zone load",
        "zone_amprion_load_mw": "Amprion zone load",
        "zone_tennet_load_mw": "TenneT zone load",
        "zone_transnetbw_load_mw": "TransnetBW zone load",
        "hour_sin": "Hour-of-day sine",
        "hour_cos": "Hour-of-day cosine",
        "dow_sin": "Day-of-week sine",
        "dow_cos": "Day-of-week cosine",
        "doy_sin": "Day-of-year sine",
        "doy_cos": "Day-of-year cosine",
        "is_weekend": "Weekend indicator",
    }
    return replacements.get(name, name.replace("_", " "))


def domain_interpretation(feature: str, sign: float) -> str:
    direction = "raises" if sign > 0 else "reduces"
    if "load" in feature and "forecast" not in feature:
        return f"Higher current or lagged load {direction} next-day high-load risk."
    if "forecast" in feature:
        return f"The external day-ahead forecast {direction} model-estimated high-load risk."
    if "renewable" in feature or "solar" in feature or "wind" in feature:
        return f"Renewable operating state {direction} next-day risk after controlling for load history."
    if "hour" in feature or "dow" in feature or "doy" in feature or "weekend" in feature:
        return f"Calendar structure {direction} high-load risk for recurring demand cycles."
    if "zone" in feature:
        return f"Regional TSO load structure {direction} national high-load risk."
    return f"The standardized feature {direction} predicted high-load risk."


def run_rq4(df: pd.DataFrame, rq1_metrics: pd.DataFrame | None = None) -> pd.DataFrame:
    X_train, X_test, y_train, y_test, _ = train_test_split_time(df, ALL_FEATURES)
    model = make_pipeline(
        StandardScaler(),
        LogisticRegression(
            max_iter=2500,
            class_weight="balanced",
            solver="lbfgs",
            random_state=RANDOM_STATE,
        ),
    )
    model.fit(X_train, y_train)
    coefs = model.named_steps["logisticregression"].coef_[0]
    coef_df = pd.DataFrame(
        {
            "feature": ALL_FEATURES,
            "display_feature": [feature_label(f) for f in ALL_FEATURES],
            "standardized_logit_coefficient": coefs,
            "absolute_coefficient": np.abs(coefs),
        }
    ).sort_values("absolute_coefficient", ascending=False)

    # A lightweight permutation check catches features that have coefficient scale
    # but little realized predictive contribution on the holdout period.
    sample = X_test.sample(n=min(3500, len(X_test)), random_state=RANDOM_STATE)
    sample_y = y_test.loc[sample.index]
    perm = permutation_importance(
        model,
        sample,
        sample_y,
        scoring="f1",
        n_repeats=6,
        random_state=RANDOM_STATE,
        n_jobs=-1,
    )
    perm_df = pd.DataFrame(
        {
            "feature": ALL_FEATURES,
            "permutation_importance_mean": perm.importances_mean,
            "permutation_importance_std": perm.importances_std,
        }
    )
    top = coef_df.head(14).merge(perm_df, on="feature", how="left")
    top["direction"] = np.where(
        top["standardized_logit_coefficient"] >= 0, "increases risk", "decreases risk"
    )
    top["domain_interpretation"] = [
        domain_interpretation(f, s)
        for f, s in zip(top["feature"], top["standardized_logit_coefficient"])
    ]
    top.to_csv(TABLE_DIR / "rq4_top_interpretable_features.csv", index=False)

    if rq1_metrics is None:
        rq1_metrics = pd.read_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv")
    transparency_scores = {
        "Logistic Regression": 5.0,
        "Linear SVM": 4.0,
        "Random Forest": 3.2,
        "Histogram Gradient Boosting": 2.8,
        "Neural Network": 1.8,
    }
    eff = rq1_metrics.copy()
    eff["transparency_score"] = eff["model"].map(transparency_scores)
    eff["f1_per_mgco2e"] = eff["f1_score"] / (
        eff["training_carbon_mgco2e_proxy"] + 1e-6
    )
    eff.to_csv(TABLE_DIR / "rq4_model_interpretability_efficiency.csv", index=False)

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11.2, 4.8), gridspec_kw={"width_ratios": [1.35, 1]})
    plot_top = top.sort_values("standardized_logit_coefficient")
    bar_colors = np.where(
        plot_top["standardized_logit_coefficient"] >= 0, "#d95f02", "#1b9aaa"
    )
    ax1.barh(
        plot_top["display_feature"],
        plot_top["standardized_logit_coefficient"],
        color=bar_colors,
        alpha=0.9,
    )
    ax1.axvline(0, color="#333333", linewidth=0.8)
    ax1.set_xlabel("Standardized logit coefficient")
    ax1.set_title("Interpretable high-load drivers")
    ax1.grid(axis="x", alpha=0.22)

    ax2.scatter(
        eff["transparency_score"],
        eff["f1_score"],
        s=100 + 0.025 * eff["model_size_kb"],
        c=eff["training_carbon_mgco2e_proxy"],
        cmap="magma_r",
        edgecolor="black",
        linewidth=0.8,
    )
    for _, row in eff.iterrows():
        ax2.annotate(
            row["model"].replace("Histogram ", "Hist. "),
            (row["transparency_score"], row["f1_score"]),
            xytext=(5, 5),
            textcoords="offset points",
            fontsize=8.3,
        )
    ax2.set_xlabel("Transparency score (higher is easier to audit)")
    ax2.set_ylabel("F1-score")
    ax2.set_title("Accuracy versus auditability")
    ax2.grid(True, alpha=0.23)
    fig.suptitle("RQ4: Interpretable Drivers and Efficiency-Auditability Trade-off")
    fig.tight_layout()
    save_figure(fig, "rq4_interpretability_efficiency")
    return top


def run_rq5(df: pd.DataFrame, rq1_metrics: pd.DataFrame | None = None) -> pd.DataFrame:
    if rq1_metrics is None:
        rq1_metrics = pd.read_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv")
    training_df = df[df["timestamp_utc"] < pd.Timestamp("2019-01-01", tz="UTC")]
    proxy = training_df["carbon_intensity_proxy_g_per_kwh"]
    scenarios = {
        "Green-window scheduling (lowest decile)": proxy[proxy <= proxy.quantile(0.10)].mean(),
        "Random-hour scheduling (training mean)": proxy.mean(),
        "Carbon-intensive scheduling (highest decile)": proxy[proxy >= proxy.quantile(0.90)].mean(),
    }
    rows = []
    for _, model_row in rq1_metrics.iterrows():
        energy_wh = model_row["energy_wh_proxy"]
        random_mg = energy_wh * scenarios["Random-hour scheduling (training mean)"]
        for scenario, scenario_proxy in scenarios.items():
            carbon_mg = energy_wh * scenario_proxy
            rows.append(
                {
                    "model": model_row["model"],
                    "scenario": scenario,
                    "training_energy_wh_proxy": energy_wh,
                    "carbon_intensity_proxy_g_per_kwh": scenario_proxy,
                    "training_carbon_mgco2e_proxy": carbon_mg,
                    "reduction_vs_random_percent": 100 * (random_mg - carbon_mg) / random_mg,
                }
            )
    result = pd.DataFrame(rows)
    result.to_csv(TABLE_DIR / "rq5_training_window_scenarios.csv", index=False)

    ordered_models = list(rq1_metrics.sort_values("f1_score", ascending=False)["model"])
    pivot = result.pivot(index="model", columns="scenario", values="training_carbon_mgco2e_proxy").loc[ordered_models]
    fig, ax = plt.subplots(figsize=(9.0, 4.9))
    x = np.arange(len(pivot))
    width = 0.25
    scenario_order = list(scenarios.keys())
    colors = ["#2a9d8f", "#457b9d", "#e76f51"]
    for i, scenario in enumerate(scenario_order):
        ax.bar(
            x + (i - 1) * width,
            pivot[scenario],
            width,
            label=scenario.split(" (")[0],
            color=colors[i],
        )
    ax.set_xticks(x)
    ax.set_xticklabels(
        [m.replace("Histogram ", "Hist. ") for m in pivot.index],
        rotation=15,
        ha="right",
    )
    ax.set_ylabel("Training carbon proxy (mg CO2e)")
    ax.set_title("RQ5: Carbon-Aware Training Window Sensitivity")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, fontsize=8.5)
    fig.tight_layout()
    save_figure(fig, "rq5_carbon_aware_training_windows")
    return result


def write_research_question_tables() -> None:
    pd.DataFrame(RESEARCH_QUESTIONS).to_csv(TABLE_DIR / "research_questions.csv", index=False)
    methodology = [
        {
            "stage": "Dataset acquisition",
            "method": "Download OPSD 60-minute single-index time-series CSV as data.csv.",
            "output": "data.csv",
        },
        {
            "stage": "Preprocessing",
            "method": "Select Germany columns, interpolate sparse gaps, engineer calendar, lag, renewable, residual-load, and TSO-zone features.",
            "output": "preprocessed_content.csv",
        },
        {
            "stage": "Target construction",
            "method": "Predict whether German load 24 hours ahead enters the top quintile of observed next-day load.",
            "output": "target_high_load_next_24h",
        },
        {
            "stage": "Model evaluation",
            "method": "Use chronological 2015-2018 training and 2019-2020 test split; compare F1, AUC, runtime, model size, and carbon proxy.",
            "output": "RQ result CSV tables and PDF figures",
        },
        {
            "stage": "Carbon-aware selection",
            "method": "Estimate runtime energy from CPU power assumptions and scale by residual-load carbon-intensity proxy.",
            "output": "Pareto frontier and scheduling scenario analysis",
        },
    ]
    pd.DataFrame(methodology).to_csv(TABLE_DIR / "methodology_map.csv", index=False)


def write_svg_icon(name: str, title: str, body: str) -> None:
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="128" height="128" viewBox="0 0 128 128" role="img" aria-label="{title}">
  <rect x="8" y="8" width="112" height="112" rx="18" fill="#ffffff" stroke="#17202a" stroke-width="4"/>
  {body}
</svg>
"""
    (ICON_DIR / name).write_text(svg, encoding="utf-8")


def write_icons() -> None:
    ICON_DIR.mkdir(parents=True, exist_ok=True)
    icons = [
        (
            "01_dataset.svg",
            "Dataset",
            '<ellipse cx="64" cy="35" rx="32" ry="12" fill="#a8dadc" stroke="#17202a" stroke-width="3"/>'
            '<path d="M32 35v46c0 7 14 12 32 12s32-5 32-12V35" fill="#dff3f5" stroke="#17202a" stroke-width="3"/>'
            '<path d="M32 58c0 7 14 12 32 12s32-5 32-12" fill="none" stroke="#17202a" stroke-width="3"/>'
            '<path d="M32 80c0 7 14 12 32 12s32-5 32-12" fill="none" stroke="#17202a" stroke-width="3"/>',
        ),
        (
            "02_research_questions.svg",
            "Research questions",
            '<circle cx="48" cy="55" r="19" fill="#ffb703" stroke="#17202a" stroke-width="3"/>'
            '<text x="48" y="63" text-anchor="middle" font-family="Arial" font-size="27" font-weight="700">?</text>'
            '<path d="M72 78l20 20" stroke="#17202a" stroke-width="8" stroke-linecap="round"/>'
            '<circle cx="52" cy="58" r="35" fill="none" stroke="#17202a" stroke-width="5"/>',
        ),
        (
            "03_preprocessing.svg",
            "Preprocessing",
            '<rect x="30" y="32" width="68" height="14" rx="7" fill="#e9ecef" stroke="#17202a" stroke-width="3"/>'
            '<rect x="30" y="57" width="68" height="14" rx="7" fill="#e9ecef" stroke="#17202a" stroke-width="3"/>'
            '<rect x="30" y="82" width="68" height="14" rx="7" fill="#e9ecef" stroke="#17202a" stroke-width="3"/>'
            '<circle cx="51" cy="39" r="9" fill="#2a9d8f"/><circle cx="77" cy="64" r="9" fill="#e76f51"/><circle cx="58" cy="89" r="9" fill="#457b9d"/>',
        ),
        (
            "04_feature_engineering.svg",
            "Feature engineering",
            '<circle cx="34" cy="64" r="12" fill="#8ecae6" stroke="#17202a" stroke-width="3"/>'
            '<circle cx="64" cy="35" r="12" fill="#ffb703" stroke="#17202a" stroke-width="3"/>'
            '<circle cx="94" cy="64" r="12" fill="#90be6d" stroke="#17202a" stroke-width="3"/>'
            '<circle cx="64" cy="94" r="12" fill="#f28482" stroke="#17202a" stroke-width="3"/>'
            '<path d="M45 57l11-13M75 44l11 13M84 71L72 86M56 86L44 71" stroke="#17202a" stroke-width="4" stroke-linecap="round"/>',
        ),
        (
            "05_model_training.svg",
            "Model training",
            '<rect x="35" y="36" width="58" height="58" rx="9" fill="#d9ed92" stroke="#17202a" stroke-width="4"/>'
            '<path d="M45 55h38M45 72h38M55 45v38M73 45v38" stroke="#17202a" stroke-width="3"/>'
            '<circle cx="46" cy="46" r="5" fill="#e76f51"/><circle cx="82" cy="82" r="5" fill="#1b9aaa"/>',
        ),
        (
            "06_carbon_accounting.svg",
            "Carbon accounting",
            '<path d="M38 87c34-2 50-25 52-52-28 4-52 20-52 52z" fill="#52b788" stroke="#17202a" stroke-width="4"/>'
            '<path d="M43 83c14-20 26-31 43-43" stroke="#ffffff" stroke-width="5" stroke-linecap="round"/>'
            '<path d="M70 28l-12 31h19L58 101l35-51H75l16-22z" fill="#ffb703" stroke="#17202a" stroke-width="3"/>',
        ),
        (
            "07_pareto_selection.svg",
            "Pareto selection",
            '<path d="M28 92H99" stroke="#17202a" stroke-width="4"/><path d="M28 92V29" stroke="#17202a" stroke-width="4"/>'
            '<circle cx="45" cy="80" r="6" fill="#e76f51"/><circle cx="60" cy="67" r="6" fill="#ffb703"/><circle cx="76" cy="55" r="6" fill="#2a9d8f"/><circle cx="91" cy="42" r="6" fill="#1b9aaa"/>'
            '<path d="M45 80C58 66 70 56 91 42" fill="none" stroke="#17202a" stroke-width="3" stroke-dasharray="5 4"/>',
        ),
        (
            "08_report_assets.svg",
            "Report assets",
            '<path d="M38 24h40l18 18v62H38z" fill="#f8f9fa" stroke="#17202a" stroke-width="4"/>'
            '<path d="M78 24v20h18" fill="none" stroke="#17202a" stroke-width="4"/>'
            '<path d="M48 58h38M48 72h38M48 86h26" stroke="#457b9d" stroke-width="5" stroke-linecap="round"/>',
        ),
    ]
    for filename, title, body in icons:
        write_svg_icon(filename, title, body)
    pd.DataFrame(
        [
            {
                "step": i + 1,
                "icon_file": item[0],
                "semantic_use": item[1],
                "format": "SVG vector",
            }
            for i, item in enumerate(icons)
        ]
    ).to_csv(ICON_DIR / "icon_catalog.csv", index=False)


def draw_workflow_icon(ax, kind: str, x: float, y: float, scale: float = 1.0) -> None:
    if kind == "database":
        ax.add_patch(Ellipse((x, y + 0.10 * scale), 0.42 * scale, 0.14 * scale, fc="#a8dadc", ec="#17202a", lw=1.2))
        ax.add_patch(Rectangle((x - 0.21 * scale, y - 0.12 * scale), 0.42 * scale, 0.22 * scale, fc="#dff3f5", ec="#17202a", lw=1.2))
        ax.add_patch(Ellipse((x, y - 0.12 * scale), 0.42 * scale, 0.14 * scale, fc="#dff3f5", ec="#17202a", lw=1.2))
    elif kind == "features":
        for dx, dy, c in [(-0.14, 0, "#8ecae6"), (0, 0.15, "#ffb703"), (0.14, 0, "#90be6d"), (0, -0.15, "#f28482")]:
            ax.add_patch(Circle((x + dx * scale, y + dy * scale), 0.06 * scale, fc=c, ec="#17202a", lw=1))
        ax.plot([x - 0.1 * scale, x, x + 0.1 * scale], [y + 0.02 * scale, y + 0.11 * scale, y + 0.02 * scale], color="#17202a", lw=1)
    elif kind == "model":
        ax.add_patch(Rectangle((x - 0.18 * scale, y - 0.16 * scale), 0.36 * scale, 0.32 * scale, fc="#d9ed92", ec="#17202a", lw=1.2))
        for off in [-0.08, 0.08]:
            ax.plot([x - 0.14 * scale, x + 0.14 * scale], [y + off * scale, y + off * scale], color="#17202a", lw=0.9)
            ax.plot([x + off * scale, x + off * scale], [y - 0.12 * scale, y + 0.12 * scale], color="#17202a", lw=0.9)
    elif kind == "leaf":
        ax.add_patch(Ellipse((x, y), 0.36 * scale, 0.22 * scale, angle=35, fc="#52b788", ec="#17202a", lw=1.2))
        ax.plot([x - 0.12 * scale, x + 0.12 * scale], [y - 0.07 * scale, y + 0.07 * scale], color="white", lw=1.4)
        ax.text(x + 0.22 * scale, y - 0.02 * scale, "CO2", ha="center", va="center", fontsize=7, fontweight="bold")
    elif kind == "chart":
        ax.plot([x - 0.18 * scale, x - 0.18 * scale, x + 0.2 * scale], [y + 0.17 * scale, y - 0.17 * scale, y - 0.17 * scale], color="#17202a", lw=1.2)
        for dx, dy, c in [(-0.08, -0.05, "#e76f51"), (0.04, 0.02, "#ffb703"), (0.15, 0.11, "#2a9d8f")]:
            ax.add_patch(Circle((x + dx * scale, y + dy * scale), 0.035 * scale, fc=c, ec="#17202a", lw=0.8))
    elif kind == "report":
        ax.add_patch(Rectangle((x - 0.15 * scale, y - 0.18 * scale), 0.30 * scale, 0.36 * scale, fc="#f8f9fa", ec="#17202a", lw=1.2))
        ax.plot([x - 0.09 * scale, x + 0.09 * scale], [y + 0.05 * scale, y + 0.05 * scale], color="#457b9d", lw=1.2)
        ax.plot([x - 0.09 * scale, x + 0.09 * scale], [y - 0.02 * scale, y - 0.02 * scale], color="#457b9d", lw=1.2)


def write_workflow_figure() -> None:
    steps = [
        ("OPSD data.csv", "German hourly load, wind, solar, TSO zones", "database"),
        ("Research questions", "Five carbon-aware ML questions", "chart"),
        ("Preprocess", "UTC time, interpolation, lags, target t+24 h", "features"),
        ("Feature families", "Calendar, load history, renewables, zones, forecast", "features"),
        ("Model suite", "LR, Linear SVM, RF, HGB, neural network", "model"),
        ("Carbon accounting", "Runtime, size, energy, residual-load proxy", "leaf"),
        ("Multi-objective evaluation", "F1, AUC, carbon, Pareto frontier", "chart"),
        ("IEEE assets", "Notebooks, figures, tables, poster deck, report", "report"),
    ]
    fig, ax = plt.subplots(figsize=(15.2, 7.5))
    ax.set_xlim(0, 15.2)
    ax.set_ylim(0, 7.5)
    ax.axis("off")
    ax.text(7.6, 7.08, "Carbon-Aware Machine Learning Model Selection Workflow", ha="center", va="center", fontsize=15.5, fontweight="bold")
    ax.text(7.6, 6.72, "Germany-aligned OPSD time series -> reproducible evidence -> IEEE/PES technical reporting", ha="center", va="center", fontsize=9.8, color="#3d405b")

    positions = [
        (1.25, 5.35),
        (3.35, 5.35),
        (5.45, 5.35),
        (7.55, 5.35),
        (10.10, 5.35),
        (12.70, 5.35),
        (12.70, 2.35),
        (7.55, 2.35),
    ]
    box_w, box_h = 1.72, 1.25
    for idx, ((title, subtitle, icon), (x, y)) in enumerate(zip(steps, positions), start=1):
        fc = "#ffffff" if idx not in (5, 6, 7) else "#f7fff7"
        box = FancyBboxPatch((x - box_w / 2, y - box_h / 2), box_w, box_h, boxstyle="round,pad=0.04,rounding_size=0.10", fc=fc, ec="#17202a", lw=1.6)
        ax.add_patch(box)
        ax.add_patch(Circle((x - box_w / 2 + 0.05, y + box_h / 2 + 0.18), 0.25, fc="#ff6b00", ec="#17202a", lw=1.3))
        ax.text(x - box_w / 2 + 0.05, y + box_h / 2 + 0.18, str(idx), ha="center", va="center", fontsize=12, color="white", fontweight="bold")
        draw_workflow_icon(ax, icon, x, y + 0.22, 1.0)
        ax.text(x, y - 0.16, title, ha="center", va="center", fontsize=10.2, fontweight="bold")
        ax.text(x, y - 0.43, textwrap.fill(subtitle, 24), ha="center", va="center", fontsize=7.9, color="#343a40")

    def arrow(start, end, color="#17202a", lw=1.7, rad=0.0, dashed=False):
        patch = FancyArrowPatch(
            start,
            end,
            arrowstyle="-|>",
            mutation_scale=16,
            lw=lw,
            color=color,
            connectionstyle=f"arc3,rad={rad}",
        )
        if dashed:
            patch.set_linestyle((0, (4, 4)))
        ax.add_patch(patch)

    for left, right in zip(positions[:5], positions[1:6]):
        arrow((left[0] + box_w / 2 + 0.07, left[1]), (right[0] - box_w / 2 - 0.07, right[1]))
    arrow((positions[5][0], positions[5][1] - box_h / 2 - 0.08), (positions[6][0], positions[6][1] + box_h / 2 + 0.08))
    arrow((positions[6][0] - box_w / 2 - 0.07, positions[6][1]), (positions[7][0] + box_w / 2 + 0.07, positions[7][1]))

    # Lightweight iteration cues: preprocessing can be revised, and model evaluation
    # loops back into training without crossing the main process path.
    arrow((positions[7][0] - 0.10, positions[7][1] + box_h / 2 + 0.04), (positions[2][0], positions[2][1] - box_h / 2 - 0.05), color="#2a9d8f", lw=1.4, rad=-0.35, dashed=True)
    arrow((positions[6][0] - 0.15, positions[6][1] + box_h / 2 + 0.04), (positions[4][0] + 0.15, positions[4][1] - box_h / 2 - 0.05), color="#e76f51", lw=1.4, rad=0.25, dashed=True)

    framework = FancyBboxPatch((8.85, 1.18), 5.55, 5.1, boxstyle="round,pad=0.05,rounding_size=0.05", fc="none", ec="#495057", lw=1.5)
    ax.add_patch(framework)
    ax.text(
        11.63,
        6.02,
        "AI evaluation framework",
        ha="center",
        va="center",
        fontsize=8.0,
        color="#495057",
        bbox={"facecolor": "white", "edgecolor": "none", "pad": 1.2},
    )
    ax.text(1.0, 0.65, "Generated assets: data.csv, preprocessed_content.csv, five notebooks, five result tables, five publication figures, reusable SVG icons, poster deck, IEEE/PES report.", fontsize=8.8, color="#3d405b")

    save_figure(fig, "workflow_carbon_aware_ml")


def add_textbox(slide, x, y, w, h, text, font_size=24, bold=False, color=(20, 30, 40), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    if align is not None:
        p.alignment = align
    return box


def add_panel(slide, x, y, w, h, title, body, fill=(248, 249, 250), accent=(42, 157, 143)):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(210, 215, 220)
    shape.line.width = Pt(1.2)
    title_box = add_textbox(slide, x + 0.25, y + 0.18, w - 0.5, 0.42, title, 18, True, accent)
    body_box = add_textbox(slide, x + 0.25, y + 0.72, w - 0.5, h - 0.9, body, 11, False, (35, 42, 50))
    body_box.text_frame.word_wrap = True
    return shape


def add_picture_if_exists(slide, path: Path, x, y, w, h):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def write_poster_deck() -> None:
    prs = Presentation()
    prs.slide_width = Inches(48)
    prs.slide_height = Inches(27)
    blank = prs.slide_layouts[6]

    title = "Accuracy versus Sustainability: Carbon-Aware ML for German Power-System Load Risk"
    subtitle = "PS26 Interdisciplinary Elective: AI, Power & Responsibility"

    # Poster option A: result-forward.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_textbox(slide, 1.2, 0.8, 45.5, 1.1, title, 34, True, (23, 32, 42), PP_ALIGN.CENTER)
    add_textbox(slide, 1.2, 2.0, 45.5, 0.6, "Poster Option A: Pareto frontier narrative", 18, False, (70, 80, 90), PP_ALIGN.CENTER)
    add_panel(slide, 1.2, 3.1, 10.7, 7.1, "Problem", "Machine-learning model selection in power systems is usually optimized for accuracy. This poster reframes the choice as a multi-objective decision: predictive performance, runtime, model size, and carbon-aware deployment cost.", fill=(241, 250, 238), accent=(42, 157, 143))
    add_panel(slide, 1.2, 10.7, 10.7, 6.6, "Dataset", "Open Power System Data hourly time series for Germany, 2015-2020. Target: whether German load 24 hours ahead enters the top quintile of next-day load.", fill=(248, 249, 250), accent=(69, 123, 157))
    add_panel(slide, 1.2, 17.8, 10.7, 7.6, "Contribution", "A reproducible carbon-aware model ranking pipeline with five research questions, publication-ready figures, CSV result tables, and IEEE/PES reporting assets.", fill=(255, 248, 230), accent=(231, 111, 81))
    add_picture_if_exists(slide, FIG_DIR / "rq1_accuracy_carbon_pareto.png", 13.0, 3.2, 20.2, 11.8)
    add_picture_if_exists(slide, FIG_DIR / "workflow_carbon_aware_ml.png", 13.0, 15.6, 20.2, 8.8)
    add_panel(slide, 34.4, 3.1, 12.3, 5.1, "Methods", "Chronological train/test split; logistic regression, linear SVM, random forest, histogram gradient boosting, and neural network; runtime and model-size instrumentation; residual-load carbon proxy.", fill=(248, 249, 250), accent=(23, 32, 42))
    add_panel(slide, 34.4, 8.8, 12.3, 5.1, "Key Evidence", "Pareto-efficient models are identified by maximizing F1-score while minimizing estimated training carbon. Seasonal testing checks whether rankings survive Germany's operating regimes.", fill=(241, 250, 238), accent=(42, 157, 143))
    add_panel(slide, 34.4, 14.5, 12.3, 5.1, "Governance Lens", "The analysis treats accuracy as one objective among several. The recommended model is the one whose marginal accuracy justifies its marginal computational and auditability burden.", fill=(255, 248, 230), accent=(231, 111, 81))
    add_panel(slide, 34.4, 20.2, 12.3, 5.2, "Takeaway", "For course presentation, lead with the Pareto plot and use the workflow as the credibility anchor: each artifact is reproducible from data.csv through the notebooks.", fill=(244, 243, 255), accent=(105, 84, 163))

    # Poster option B: five-RQ evidence matrix.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(252, 253, 255)
    add_textbox(slide, 1.2, 0.75, 45.5, 1.05, title, 32, True, (23, 32, 42), PP_ALIGN.CENTER)
    add_textbox(slide, 1.2, 1.9, 45.5, 0.6, "Poster Option B: five research questions as an evidence matrix", 18, False, (70, 80, 90), PP_ALIGN.CENTER)
    rq_figs = [
        ("RQ1", "Accuracy-carbon Pareto", "rq1_accuracy_carbon_pareto.png"),
        ("RQ2", "Feature value", "rq2_feature_ablation.png"),
        ("RQ3", "Seasonal robustness", "rq3_seasonal_f1_heatmap.png"),
        ("RQ4", "Interpretability", "rq4_interpretability_efficiency.png"),
        ("RQ5", "Training windows", "rq5_carbon_aware_training_windows.png"),
    ]
    x_positions = [1.2, 17.0, 32.8, 1.2, 24.9]
    y_positions = [3.2, 3.2, 3.2, 15.0, 15.0]
    widths = [14.6, 14.6, 14.0, 21.9, 21.8]
    heights = [10.8, 10.8, 10.8, 10.0, 10.0]
    for (rq, caption, fig_name), x, y, w, h in zip(rq_figs, x_positions, y_positions, widths, heights):
        add_panel(slide, x, y, w, h, f"{rq}: {caption}", "", fill=(255, 255, 255), accent=(42, 157, 143))
        add_picture_if_exists(slide, FIG_DIR / fig_name, x + 0.55, y + 1.05, w - 1.1, h - 1.65)
    add_textbox(slide, 1.2, 25.45, 45.5, 0.55, subtitle, 14, False, (70, 80, 90), PP_ALIGN.CENTER)

    # Poster option C: governance/recommendations.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_textbox(slide, 1.2, 0.8, 45.5, 1.1, title, 33, True, (23, 32, 42), PP_ALIGN.CENTER)
    add_textbox(slide, 1.2, 2.0, 45.5, 0.6, "Poster Option C: responsible AI and deployment governance briefing", 18, False, (70, 80, 90), PP_ALIGN.CENTER)
    add_picture_if_exists(slide, FIG_DIR / "workflow_carbon_aware_ml.png", 1.4, 3.2, 28.8, 14.3)
    add_panel(slide, 31.2, 3.2, 15.4, 4.5, "Decision Principle", "Select models on a Pareto frontier: accuracy, carbon proxy, memory footprint, inference latency, and auditability must be reviewed together.", fill=(241, 250, 238), accent=(42, 157, 143))
    add_panel(slide, 31.2, 8.3, 15.4, 4.5, "Operational Policy", "Schedule retraining in lower residual-load windows where practicable; publish model cards that include runtime, model size, and carbon-proxy assumptions.", fill=(248, 249, 250), accent=(69, 123, 157))
    add_panel(slide, 31.2, 13.4, 15.4, 4.5, "Scientific Control", "Use chronological splits, seasonal stress tests, and feature ablations to separate real predictive value from coincidental correlations.", fill=(255, 248, 230), accent=(231, 111, 81))
    add_panel(slide, 1.4, 18.3, 14.2, 6.9, "Primary Result", "The full report ranks models by carbon-adjusted score and marks Pareto-efficient candidates. This makes the final recommendation traceable rather than preference-driven.", fill=(244, 243, 255), accent=(105, 84, 163))
    add_panel(slide, 16.8, 18.3, 14.2, 6.9, "Limitations", "The carbon estimate is a proxy based on runtime, assumed CPU power, and residual-load intensity. It should be replaced with measured hardware energy for formal deployment.", fill=(248, 249, 250), accent=(23, 32, 42))
    add_panel(slide, 32.2, 18.3, 14.4, 6.9, "Recommended Use", "Use this poster for a methods-and-governance oral presentation: start with the workflow, then defend the model recommendation using RQ1 and RQ5.", fill=(241, 250, 238), accent=(42, 157, 143))

    prs.save(POSTER_DIR / "carbon_aware_ml_poster_options.pptx")


def notebook_json(title: str, rq: dict, runner: str) -> dict:
    md = f"""# {title}

**Research question.** {rq['question']}

This notebook is intentionally reproducible: it reads the raw `data.csv` if `preprocessed_content.csv` is absent, rebuilds the engineered dataset, and writes the publication-ready table and figure for {rq['id']}.

Methodological steps:

1. Load the Germany-relevant OPSD hourly time-series variables.
2. Construct a next-day high-load target.
3. Use chronological evaluation: 2015-2018 for training and 2019-2020 for testing.
4. Save the result table as `outputs/tables/{rq['table']}`.
5. Save the figure as `outputs/figures/{rq['figure']}` and a PNG companion for the report.
"""
    code = f"""import os
import sys
from pathlib import Path

PROJECT_ROOT = Path.cwd().parent if Path.cwd().name == "notebooks" else Path.cwd()
sys.path.insert(0, str(PROJECT_ROOT))

from scripts import generate_analysis_assets as assets

df = assets.load_or_create_preprocessed()
result = assets.{runner}(df)
result.head()
"""
    return {
        "cells": [
            {
                "cell_type": "markdown",
                "metadata": {},
                "source": md.splitlines(keepends=True),
            },
            {
                "cell_type": "code",
                "execution_count": None,
                "metadata": {},
                "outputs": [],
                "source": code.splitlines(keepends=True),
            },
        ],
        "metadata": {
            "kernelspec": {
                "display_name": "Python 3",
                "language": "python",
                "name": "python3",
            },
            "language_info": {"name": "python", "pygments_lexer": "ipython3"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }


def write_notebooks() -> None:
    runners = ["run_rq1", "run_rq2", "run_rq3", "run_rq4", "run_rq5"]
    for rq, runner in zip(RESEARCH_QUESTIONS, runners):
        title = f"{rq['id']} - Carbon-Aware ML Analysis"
        path = NOTEBOOK_DIR / f"{rq['id'].lower()}_carbon_aware_ml.ipynb"
        path.write_text(
            json.dumps(notebook_json(title, rq, runner), indent=2),
            encoding="utf-8",
        )


def write_readme() -> None:
    readme = f"""# Carbon-Aware Machine Learning Model Selection for Germany

This project package was generated for the PS26 final project prompt.

## Core Files

- `data.csv`: raw OPSD hourly power-system time-series file.
- `preprocessed_content.csv`: Germany-focused engineered modeling table.
- `notebooks/`: five runnable notebooks, one per research question.
- `outputs/tables/`: CSV result tables and project metadata.
- `outputs/figures/`: publication figures as PDF plus PNG companions for documents and posters.
- `icons/`: reusable SVG icons used in the workflow figure.
- `posters/carbon_aware_ml_poster_options.pptx`: three editable PowerPoint poster concepts.
- `report/Carbon_Aware_ML_Germany_IEEEPES_Technical_Report.docx`: IEEE/PES-style technical report.

## Research Questions

{chr(10).join(f"- **{rq['id']}**: {rq['question']}" for rq in RESEARCH_QUESTIONS)}

## Reproduce

Run:

```bash
python3 scripts/generate_analysis_assets.py
```

The notebooks call the same generation functions and can be run independently.
"""
    (ROOT / "README.md").write_text(readme, encoding="utf-8")


def run_all() -> None:
    ensure_dirs()
    df = load_or_create_preprocessed(force=True)
    write_research_question_tables()
    rq1_metrics, fitted, test_df, y_test = run_rq1(df, return_artifacts=True)
    run_rq2(df)
    run_rq3(df, fitted, test_df, y_test)
    run_rq4(df, rq1_metrics)
    run_rq5(df, rq1_metrics)
    write_icons()
    write_workflow_figure()
    write_poster_deck()
    write_notebooks()
    write_readme()


if __name__ == "__main__":
    run_all()
