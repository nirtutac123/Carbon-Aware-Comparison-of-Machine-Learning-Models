from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "Carbon_Aware_ML_Project_Documentation.docx"
FIG_DIR = ROOT / "outputs" / "figures"
TABLE_DIR = ROOT / "outputs" / "tables"


def set_run(run, size=10.5, bold=False, italic=False, color=(30, 42, 55)):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", size=10.5, bold=False, color=(30, 42, 55), align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, size=15 if level == 1 else 12.5, bold=True, color=(15, 55, 78))
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, size=8.7, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    set_run(r, size=size, bold=bold, color=(25, 39, 52))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, title, rows, headers):
    add_para(doc, title, size=10.2, bold=True, color=(15, 55, 78), align=WD_ALIGN_PARAGRAPH.CENTER)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    for j, header in enumerate(headers):
        shade_cell(table.rows[0].cells[j], "DCEAF7")
        set_cell(table.rows[0].cells[j], header, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell(cells[j], value, size=8.2, align=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", size=3)
    return table


def add_picture(doc, filename, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(FIG_DIR / filename), width=Inches(width))
    inline._inline.docPr.set("title", caption.split(". ", 1)[0])
    inline._inline.docPr.set("descr", caption)
    add_para(doc, caption, size=8.8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def configure(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.header.paragraphs[0].text = "Project Documentation | Carbon-Aware AI Model Selection"
        section.footer.paragraphs[0].text = "PS26 reproducibility documentation | May 2026"
        section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for p in list(section.header.paragraphs) + list(section.footer.paragraphs):
            for r in p.runs:
                set_run(r, size=8.5, color=(90, 100, 108))


def main():
    doc = Document()
    configure(doc)
    metadata = json.loads((TABLE_DIR / "dataset_metadata.json").read_text(encoding="utf-8"))
    rq1 = pd.read_csv(TABLE_DIR / "rq1_model_tradeoff_metrics.csv")
    rq2 = pd.read_csv(TABLE_DIR / "rq2_feature_ablation_metrics.csv")
    rq5 = pd.read_csv(TABLE_DIR / "rq5_training_window_scenarios.csv")
    best = rq1.sort_values("f1_score", ascending=False).iloc[0]
    low = rq1.sort_values("training_carbon_mgco2e_proxy").iloc[0]
    green_reduction = rq5[rq5["scenario"].str.startswith("Green")]["reduction_vs_random_percent"].iloc[0]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("How the Carbon-Aware AI Model Selection Project Was Built")
    set_run(r, size=20, bold=True, color=(15, 55, 78))
    add_para(doc, "A reproducibility and methodology companion for the PS26 project", size=12.5, bold=True, color=(70, 86, 96), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "This document explains the complete workflow used to create the data assets, notebooks, figures, tables, workflow diagram, posters, and IEEE/PES-style technical report.", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_picture(doc, "workflow_structured_ai_esg_style.png", "Fig. 1. Structured workflow used to build and audit the project assets.", width=6.6)

    add_heading(doc, "1. Project Objective", 1)
    add_para(doc, "The project studies model selection for German power-system analytics as a responsible AI problem. Instead of ranking machine-learning models only by predictive accuracy, the analysis compares accuracy, F1-score, ROC-AUC, runtime, serialized model size, and a transparent carbon proxy. The final empirical task predicts whether German electricity load 24 hours ahead enters the highest quintile of observed demand.")

    add_heading(doc, "2. Data Source and Access Link", 1)
    add_para(doc, "The raw dataset is Open Power System Data (OPSD) Time Series, version 2020-10-06. It contains load, wind and solar generation, prices, and capacities in hourly resolution, with variables aggregated by country, bidding zone, or control area. The project uses Germany-specific load, forecast, wind, solar, and TSO-zone variables.")
    add_table(
        doc,
        "TABLE 1. Data source link and local project files.",
        [
            ["Direct CSV link", "https://data.open-power-system-data.org/time_series/2020-10-06/time_series_60min_singleindex.csv"],
            ["Metadata link", "https://data.open-power-system-data.org/time_series/2020-10-06/datapackage.json"],
            ["Local raw file", "data.csv"],
            ["Local processed file", "preprocessed_content.csv"],
            ["Processed observations", f"{metadata['n_rows']:,} rows"],
            ["Engineered features", f"{metadata['n_features']} features"],
        ],
        ["Item", "Value"],
    )

    add_heading(doc, "3. Research Design", 1)
    add_para(doc, "Five research questions were defined to match the course prompt and the IEEE-style technical-report framing: model trade-offs, feature value, seasonal robustness, interpretability, and carbon-aware training windows. Each research question has a separate notebook, a publication-ready PDF figure, and a CSV table.")
    rq_rows = pd.read_csv(TABLE_DIR / "research_questions.csv")[["id", "question", "figure", "table"]].values.tolist()
    add_table(doc, "TABLE 2. Research questions and generated artifacts.", rq_rows, ["ID", "Research question", "Figure", "Table"])

    add_heading(doc, "4. Data Preparation", 1)
    add_para(doc, "The preprocessing script reads the raw OPSD CSV, selects Germany-relevant variables, interpolates sparse gaps, sorts observations by UTC timestamp, and constructs a chronological modeling table. The target is `target_high_load_next_24h`, equal to 1 when German load 24 hours ahead is at or above the 80th percentile of next-day load.")
    add_table(
        doc,
        "TABLE 3. Main engineered feature families.",
        [
            ["Calendar", "hour, day-of-week, day-of-year cyclic encodings; weekend flag"],
            ["Load history", "current load; 1 h, 24 h, and 168 h lags; rolling means and volatility"],
            ["Renewables", "solar, wind, renewable share, residual load, residual-load share"],
            ["Regional/TSO", "50Hertz, Amprion, TenneT, TransnetBW, and zone imbalance"],
            ["Forecast", "German load forecast from OPSD/ENTSO-E transparency source"],
        ],
        ["Feature family", "Variables"],
    )

    add_heading(doc, "5. Modeling and Evaluation", 1)
    add_para(doc, "The model suite compares logistic regression, linear SVM, random forest, histogram gradient boosting, and a shallow neural network. The split is chronological: 2015-2018 observations are used for training, while 2019-2020 observations are used for testing. The boosted-tree comparator uses scikit-learn histogram gradient boosting so the package remains reproducible without an external XGBoost dependency.")
    add_table(
        doc,
        "TABLE 4. Key model results.",
        rq1[["model", "f1_score", "roc_auc", "training_carbon_mgco2e_proxy", "pareto_efficient"]].round(3).values.tolist(),
        ["Model", "F1", "ROC-AUC", "Carbon proxy (mg CO2e)", "Pareto"],
    )
    add_picture(doc, "rq1_accuracy_carbon_pareto.png", "Fig. 2. Accuracy-carbon Pareto result used in the report and posters.", width=5.8)

    add_heading(doc, "6. Carbon-Aware Accounting", 1)
    add_para(doc, "The carbon estimate is a proxy. Training and prediction energy are estimated from measured runtime and assumed CPU power values. The carbon intensity proxy scales with residual load, meaning hours with lower wind/solar coverage receive higher proxy intensity. This is appropriate for comparative sensitivity analysis but should be replaced with measured hardware power and official emission factors in a deployment study.")
    add_para(doc, f"The strongest model in this run is {best['model']} with F1 = {best['f1_score']:.3f}; the lowest-compute model is {low['model']} with a carbon proxy of {low['training_carbon_mgco2e_proxy']:.3f} mg CO2e. Green-window scheduling reduces the training-carbon proxy by {green_reduction:.1f}% relative to random-hour scheduling.")

    add_heading(doc, "7. Generated Assets", 1)
    add_table(
        doc,
        "TABLE 5. Project asset map.",
        [
            ["Raw data", "data.csv"],
            ["Processed data", "preprocessed_content.csv"],
            ["Notebooks", "notebooks/rq1_carbon_aware_ml.ipynb through rq5_carbon_aware_ml.ipynb"],
            ["Tables", "outputs/tables/*.csv"],
            ["Figures", "outputs/figures/*.pdf and *.png"],
            ["Workflow", "outputs/figures/workflow_structured_ai_esg_style.pdf/png"],
            ["Posters", "posters/carbon_aware_ml_academic_posters_v2.pptx"],
            ["Technical report", "report/Carbon_Aware_ML_Germany_IEEEPES_Technical_Report.docx"],
        ],
        ["Asset type", "Location"],
    )

    add_heading(doc, "8. Poster and Workflow Design", 1)
    add_para(doc, "The revised poster deck follows the provided lecture screenshot more closely: a centered title, top corner icons, upper Motivation and Datasets panels, a full-height left methodology strip, and right-side sections for research questions, key results, findings, and conclusion. Three visual variants are included in one editable PowerPoint deck. The structured workflow diagram follows an academic 1-8 left-to-right flow with circular step markers, an AI framework box, and explicit outputs and recommendations.")

    add_heading(doc, "9. IEEE/PES Report Verification", 1)
    add_para(doc, "The technical report was generated from the provided IEEE PES Technical Report template rather than from a blank document. The final report keeps the expected front matter pattern (cover, acknowledgments, keywords, contents), numbered technical sections, centered table labels, centered figure captions, references, and a restrained Times-style technical-report layout. The report was rendered to page images and visually inspected for clipping, orphaned captions, table crowding, and broken figure placement. It also passed the document accessibility audit after adding figure alt text and repeated table-header metadata.")
    add_table(
        doc,
        "TABLE 6. Report-template compliance checklist.",
        [
            ["Template base", "Built from PES-Technical-Report-Template_Jan_2019.docx", "Pass"],
            ["Front matter", "Acknowledgments, keywords, contents", "Pass"],
            ["Technical sections", "Introduction, dataset, methodology, results, governance, limitations, conclusions", "Pass"],
            ["Tables/figures", "Centered labels/captions and placed near citations", "Pass"],
            ["References", "Numbered technical-report style references", "Pass"],
            ["Visual QA", "Rendered with artifact-tool and inspected", "Pass"],
            ["Accessibility", "0 high, 0 medium, 0 low findings", "Pass"],
        ],
        ["Check", "Evidence", "Status"],
    )

    add_heading(doc, "10. How to Regenerate", 1)
    add_para(doc, "From the project root, run `python3 scripts/generate_analysis_assets.py` to rebuild the data, notebooks, tables, and core figures. Run `python3 scripts/generate_revised_poster_and_workflow.py` to regenerate the revised workflow and poster deck. Run the bundled Python command for `scripts/generate_report_docx.py` and `scripts/generate_project_documentation_docx.py` to rebuild the Word deliverables.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
